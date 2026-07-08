"""Erzeugt das Steuerungs-Template 'Kapazitaets-Allokationsmatrix' fuer das
Capacity-Board (Homebase x Value Stream).

Editierbares Excel mit Live-Formeln:
- Blatt 'Allokation %': je Homebase die %-Verteilung auf Value Streams / Chapter /
  Run & Line; Zeilensumme muss 100 % ergeben (Ampel-Check).
- Blatt 'FTE & Deckung': rechnet aus % x FTE das Angebot je Value Stream, stellt
  es der geplanten Nachfrage gegenueber und zeigt die Deckungsluecke (Gap).
Werte sind Beispielwerte und im Board zu befuellen.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

DST = "2026_07_08_Kapazitaets_Allokationsmatrix.xlsx"

# ---- Hausfarben ----
ORANGE = "F07F12"; DARK = "1F2937"; HOMEBASE = "2E5A8F"; LIGHT = "F7F7F7"
ROW_ALT = "EEF1F4"; GREEN = "2E7D32"; RED = "B91C1C"; WHITE = "FFFFFF"
GREY = "4B5563"; AMBER = "B45309"

thin = Side(style="thin", color="D1D5DB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(hex_): return PatternFill("solid", fgColor=hex_)
def fnt(**k): return Font(name="Calibri", **k)

# ---- Daten (Beispielwerte) ----
homebases = [
    ("Workplace Solutions", 18),
    ("Infrastructure & Platforms", 30),
    ("Local IT Operations", 22),
    ("Business Solutions (Design/AMS)", 26),
    ("Application Technology (SAP Basis/Dev)", 14),
    ("Cyber Defense / SecOps", 12),
    ("Cyber GRC", 8),
    ("EA & Innovation", 6),
]
streams = [
    "Modern Workplace", "Network & Connectivity", "Compute/Cloud/DC",
    "Identity & Access", "SAP & Business Apps", "Data/Analytics/KI",
    "Service Desk / ITSM", "Chapter / CoP", "Run & Line",
]
# %-Allokation je Homebase (Reihenfolge = streams), Summe je Zeile = 100
alloc = {
    "Workplace Solutions":                    [70, 0, 0, 0, 0, 0, 5, 15, 10],
    "Infrastructure & Platforms":             [0, 25, 30, 20, 0, 0, 0, 10, 15],
    "Local IT Operations":                    [30, 10, 0, 0, 0, 0, 40, 5, 15],
    "Business Solutions (Design/AMS)":        [0, 0, 0, 5, 40, 25, 0, 15, 15],
    "Application Technology (SAP Basis/Dev)": [0, 0, 0, 0, 60, 10, 0, 15, 15],
    "Cyber Defense / SecOps":                 [5, 5, 5, 10, 0, 0, 5, 20, 50],
    "Cyber GRC":                              [0, 0, 0, 10, 0, 0, 0, 30, 60],
    "EA & Innovation":                        [0, 0, 0, 0, 0, 20, 0, 40, 40],
}
# geplante Nachfrage je Ziel (FTE) - unabhaengiger Planungsinput
demand = [22, 10, 10, 9, 18, 10, 10, None, None]  # Chapter/Run&Line ohne Ziel-Gap

wb = Workbook()

# ======================================================================
# BLATT 1  --  Allokation %
# ======================================================================
ws = wb.active
ws.title = "Allokation %"
ws.sheet_view.showGridLines = False

ws["A1"] = "Kapazitäts-Allokationsmatrix — Homebase × Value Stream"
ws["A1"].font = fnt(bold=True, size=15, color=DARK)
ws["A2"] = "Steuerungstemplate Capacity-Board · %-Verteilung der Homebase-Kapazität · Beispielwerte"
ws["A2"].font = fnt(size=10, italic=True, color=GREY)

HDR = 4  # Kopfzeile
# Spalten: A Homebase | B FTE gesamt | C..K streams | L Summe %
ws.cell(HDR, 1, "Homebase").font = fnt(bold=True, color=WHITE)
ws.cell(HDR, 2, "FTE gesamt").font = fnt(bold=True, color=WHITE)
for j, s in enumerate(streams):
    c = ws.cell(HDR, 3 + j, s)
    c.font = fnt(bold=True, color=WHITE, size=9)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
sum_col = 3 + len(streams)
ws.cell(HDR, sum_col, "Summe %").font = fnt(bold=True, color=WHITE)
for col in range(1, sum_col + 1):
    cc = ws.cell(HDR, col)
    cc.fill = fill(DARK); cc.border = box
    if col in (2, sum_col) or col >= 3:
        cc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

first_data = HDR + 1
for i, (hb, fte) in enumerate(homebases):
    r = first_data + i
    a = ws.cell(r, 1, hb); a.font = fnt(bold=True, color=DARK, size=10); a.border = box
    b = ws.cell(r, 2, fte); b.font = fnt(color=HOMEBASE, bold=True); b.border = box
    b.alignment = Alignment(horizontal="center")
    for j, v in enumerate(alloc[hb]):
        cell = ws.cell(r, 3 + j, v)
        cell.number_format = '0"%"'
        cell.alignment = Alignment(horizontal="center")
        cell.border = box
        if v == 0:
            cell.font = fnt(color="C7CDD4", size=10)
        else:
            cell.font = fnt(color=DARK, size=10)
    # Summe-Formel
    c1 = get_column_letter(3); c2 = get_column_letter(2 + len(streams))
    sc = ws.cell(r, sum_col, f"=SUM({c1}{r}:{c2}{r})")
    sc.number_format = '0"%"'; sc.font = fnt(bold=True); sc.border = box
    sc.alignment = Alignment(horizontal="center")
    # Zeilen-Zebra
    if i % 2 == 1:
        for col in range(1, sum_col + 1):
            if ws.cell(r, col).fill.fgColor.rgb in (None, "00000000"):
                ws.cell(r, col).fill = fill(ROW_ALT)

last_data = first_data + len(homebases) - 1
# Summenzeile FTE
tr = last_data + 1
ws.cell(tr, 1, "Σ FTE-Basis").font = fnt(bold=True, color=DARK)
tcell = ws.cell(tr, 2, f"=SUM(B{first_data}:B{last_data})")
tcell.font = fnt(bold=True, color=HOMEBASE); tcell.alignment = Alignment(horizontal="center")
for col in range(1, 3):
    ws.cell(tr, col).border = box; ws.cell(tr, col).fill = fill(LIGHT)

# Ampel: Summe % != 100 -> rot, ==100 -> gruen
sum_letter = get_column_letter(sum_col)
rng = f"{sum_letter}{first_data}:{sum_letter}{last_data}"
ws.conditional_formatting.add(rng, CellIsRule(operator="notEqual", formula=["100"],
    fill=fill("FDE2E2"), font=fnt(bold=True, color=RED)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["100"],
    fill=fill("E3F2EA"), font=fnt(bold=True, color=GREEN)))

# Spaltenbreiten
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 11
for j in range(len(streams)):
    ws.column_dimensions[get_column_letter(3 + j)].width = 12
ws.column_dimensions[get_column_letter(sum_col)].width = 10
ws.row_dimensions[HDR].height = 46

note_r = tr + 2
ws.cell(note_r, 1, "Regel: Jede Homebase-Zeile muss 100 % ergeben (grün). "
        "Rot = Über-/Unterallokation. Werte im Board pflegen.").font = fnt(italic=True, size=9, color=GREY)

# ======================================================================
# BLATT 2  --  FTE & Deckung
# ======================================================================
ws2 = wb.create_sheet("FTE & Deckung")
ws2.sheet_view.showGridLines = False
ws2["A1"] = "FTE-Angebot vs. Nachfrage je Value Stream"
ws2["A1"].font = fnt(bold=True, size=15, color=DARK)
ws2["A2"] = "Angebot = Σ (Homebase-FTE × %-Allokation). Nachfrage = Planungsinput. Gap = Angebot − Nachfrage."
ws2["A2"].font = fnt(size=10, italic=True, color=GREY)

H2 = 4
heads2 = ["Ziel (Value Stream / Sammelposten)", "Angebot (FTE)", "Nachfrage (FTE)", "Deckung (Gap)"]
for j, h in enumerate(heads2):
    c = ws2.cell(H2, 1 + j, h)
    c.font = fnt(bold=True, color=WHITE); c.fill = fill(ORANGE); c.border = box
    c.alignment = Alignment(horizontal="center" if j else "left", vertical="center", wrap_text=True)

for i, s in enumerate(streams):
    r = H2 + 1 + i
    col_letter = get_column_letter(3 + i)  # Spalte im Blatt 1
    nm = ws2.cell(r, 1, s); nm.font = fnt(bold=True, color=DARK, size=10); nm.border = box
    # Angebot = SUMPRODUCT(FTE, %-Spalte)/100
    supply = (f"=SUMPRODUCT('Allokation %'!$B${first_data}:$B${last_data},"
              f"'Allokation %'!{col_letter}{first_data}:{col_letter}{last_data})/100")
    sc = ws2.cell(r, 2, supply); sc.number_format = "0.0"
    sc.alignment = Alignment(horizontal="center"); sc.border = box; sc.font = fnt(color=HOMEBASE, bold=True)
    # Nachfrage
    dem = demand[i]
    dc = ws2.cell(r, 3, dem if dem is not None else "—")
    dc.number_format = "0.0"; dc.alignment = Alignment(horizontal="center"); dc.border = box
    dc.font = fnt(color=DARK)
    # Gap
    if dem is not None:
        gc = ws2.cell(r, 4, f"=B{r}-C{r}")
        gc.number_format = "+0.0;-0.0"
    else:
        gc = ws2.cell(r, 4, "—")
    gc.alignment = Alignment(horizontal="center"); gc.border = box; gc.font = fnt(bold=True)
    if i % 2 == 1:
        for col in range(1, 5):
            ws2.cell(r, col).fill = fill(ROW_ALT)

last2 = H2 + len(streams)
# Gap-Ampel
gap_rng = f"D{H2+1}:D{last2}"
ws2.conditional_formatting.add(gap_rng, CellIsRule(operator="lessThan", formula=["0"],
    fill=fill("FDE2E2"), font=fnt(bold=True, color=RED)))
ws2.conditional_formatting.add(gap_rng, CellIsRule(operator="greaterThanOrEqual", formula=["0"],
    fill=fill("E3F2EA"), font=fnt(bold=True, color=GREEN)))

ws2.column_dimensions["A"].width = 40
for col in "BCD":
    ws2.column_dimensions[col].width = 16
ws2.row_dimensions[H2].height = 30

ws2.cell(last2 + 2, 1, "Rot = Unterdeckung (Nachfrage > Angebot) → Priorisierung/Nachschärfung im "
         "Capacity-Board. Chapter & Run & Line ohne Ziel-Gap (keine Wertstrom-Nachfrage).").font = \
    fnt(italic=True, size=9, color=GREY)

# ======================================================================
# BLATT 3  --  Anleitung
# ======================================================================
ws3 = wb.create_sheet("Anleitung", 0)
ws3.sheet_view.showGridLines = False
ws3.column_dimensions["A"].width = 3
ws3.column_dimensions["B"].width = 100
ws3["B2"] = "Kapazitäts-Allokationsmatrix — Anleitung"
ws3["B2"].font = fnt(bold=True, size=16, color=DARK)
lines = [
    ("Zweck", True),
    ("Steuerungstemplate für das Capacity-Board. Macht sichtbar, wie viel Kapazität jede "
     "Homebase (Linie) an welche Mannschaft (Value Stream) zusagt — und ob die Summe die "
     "geplante Nachfrage deckt.", False),
    ("", False),
    ("So wird es genutzt", True),
    ("1.  Blatt 'Allokation %': FTE-Basis je Homebase eintragen, dann die %-Zusage je "
     "Value Stream / Chapter / Run & Line pflegen. Jede Zeile MUSS 100 % ergeben (Spalte "
     "'Summe %' wird grün; rot = Fehler).", False),
    ("2.  Blatt 'FTE & Deckung': rechnet automatisch das FTE-Angebot je Value Stream und "
     "vergleicht es mit der Nachfrage. 'Deckung (Gap)' rot = Unterdeckung.", False),
    ("3.  Bei roter Deckung: im Board entscheiden — Nachfrage priorisieren, Kapazität "
     "umschichten, Dienstleister zuschalten oder Skill aufbauen (Chapter).", False),
    ("", False),
    ("Spielregeln (aus dem Zielmodell)", True),
    ("•  Kapazität wird ZUGESAGT, nicht angenommen — die %-Werte sind ein Commitment der "
     "Homebase.", False),
    ("•  Eine Priorität pro Person; Dediziert vor geteilt (Splitting über > 2 Streams "
     "vermeiden).", False),
    ("•  Homebase ist accountable für die Kapazitäts-Zusage, Service Owner für die "
     "Priorität. Ressourcenkonflikte entscheidet das Capacity-Board.", False),
    ("", False),
    ("Rhythmus", True),
    ("Quartalsweise Grundallokation (Portfolio-Board), monatliche Nachjustierung (Service "
     "Management Review). Alle Werte hier sind Beispielwerte.", False),
]
r = 4
for text, is_h in lines:
    c = ws3.cell(r, 2, text)
    c.font = fnt(bold=is_h, size=12 if is_h else 10.5, color=ORANGE if is_h else DARK)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.row_dimensions[r].height = 30 if (is_h or len(text) > 90) else 16
    r += 1

wb.save(DST)
print("saved", DST)
