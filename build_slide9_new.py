"""Insert a new slide right after slide 8: ServiceNow – Plattform, SPM Pro & STIHL-Kontext.

Sections:
1) Was bietet ServiceNow als Tool? (Now Platform)
2) Was leistet SPM Pro?
3) STIHL-Kontext: seit 2014 global, Hosting Frankfurt/Düsseldorf DSGVO
4) Deutsche Referenzkunden (Logo-Platzhalter mit Firmennamen)
"""
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v4.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v5.pptx"

ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)

U = 76200  # EMU per "unit" – 160x90 grid

def ex(u): return int(u * U)

prs = Presentation(SRC)

# pick a blank layout – use the same layout as slide 8
layout = prs.slides[7].slide_layout
new_slide = prs.slides.add_slide(layout)

# wipe placeholders inherited from layout
for shp in list(new_slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

slide = new_slide

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.2)
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK, default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(50000); tf.margin_right = Emu(50000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)
        if "space_after" in st:
            p.space_after = Pt(st["space_after"])

def add_text(left, top, w, h, paragraphs, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = kw.pop("anchor", MSO_ANCHOR.TOP)
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ============== HEADER ==============
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("ServiceNow – Plattform, SPM Pro und STIHL-Kontext",
     dict(size=22, bold=True, color=WHITE)),
])
add_text(3, 3.2, 140, 2.4, [
    ("Industriestandard, Pro-Modul für Strategic Portfolio Management und bestehende STIHL-Basis",
     dict(size=12, italic=True, color=LIGHT_GREY)),
])
add_rect(150, 0, 10, 6, ORANGE)
tb = slide.shapes.add_textbox(ex(150), ex(0), ex(10), ex(6))
tb.fill.background(); tb.line.fill.background()
tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_paragraphs(tb, [("Folie 9", dict(size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER))])

# ============== 3 COLUMNS ==============
col_top = 8
col_h = 38
gap = 1.5
total = 152
cw = (total - 2 * gap) / 3

# --- Column 1: Now Platform ---
x1 = 4
add_rect(x1, col_top, cw, 3.8, ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(x1, col_top, cw, 3.8, [
    ("ServiceNow Now Platform",
     dict(size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
], anchor=MSO_ANCHOR.MIDDLE)
add_rect(x1, col_top + 4, cw, col_h - 4, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

c1_paragraphs = [
    ("Cloud-native Workflow-Plattform – eine Datenbasis für IT, Mitarbeiter, Kunden, Risiko",
     dict(size=10.5, bold=True, color=DARK, space_after=6)),
    ("Produktbreite (Auswahl):", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  IT Service Management (ITSM) & IT Operations (ITOM)", dict(size=9.5)),
    ("•  Strategic Portfolio Management (SPM)", dict(size=9.5)),
    ("•  HR Service Delivery & Employee Experience", dict(size=9.5)),
    ("•  Customer Service Management (CSM)", dict(size=9.5)),
    ("•  Security Operations & Governance, Risk, Compliance",
     dict(size=9.5, space_after=8)),
    ("Plattform-Bausteine:", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  Workflow Engine, Now Assist (KI), Performance Analytics", dict(size=9.5)),
    ("•  App Engine / Low-Code Studio – kundenspezifische Apps ohne Eigenentwicklung",
     dict(size=9.5)),
    ("•  Integration Hub – Standard-Konnektoren zu SAP, AD, Cloud-Diensten",
     dict(size=9.5, space_after=8)),
    ("> 85 % der Fortune 500 nutzen ServiceNow",
     dict(size=10, bold=True, color=ORANGE)),
]
add_text(x1 + 0.8, col_top + 4.6, cw - 1.6, col_h - 5, c1_paragraphs)

# --- Column 2: SPM Pro ---
x2 = x1 + cw + gap
add_rect(x2, col_top, cw, 3.8, ORANGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(x2, col_top, cw, 3.8, [
    ("ServiceNow SPM Pro",
     dict(size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
], anchor=MSO_ANCHOR.MIDDLE)
add_rect(x2, col_top + 4, cw, col_h - 4, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

c2_paragraphs = [
    ("End-to-end von der Strategie bis zur Umsetzung – auf einer Plattform",
     dict(size=10.5, bold=True, color=DARK, space_after=6)),
    ("Strategic Planning & Goal Framework (OKR)", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  Strategien, Ziele, Investitionsrahmen verknüpft mit Portfolio",
     dict(size=9.5, space_after=6)),
    ("Demand Management", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  Zentrale Erfassung, Bewertung und Priorisierung aller Demands",
     dict(size=9.5, space_after=6)),
    ("Project Portfolio Management (PPM)", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  Klassisch, agil, hybrid – Gantt, Kanban, SAFe-Workflows", dict(size=9.5, space_after=6)),
    ("Resource & Financial Management", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  Ressourcen-, Kapazitäts-, Budget- und Szenarioplanung (What-if)",
     dict(size=9.5, space_after=6)),
    ("Application Portfolio & Benefits", dict(size=10, bold=True, color=DARK, space_after=2)),
    ("•  Anwendungsportfolio, Nutzen-Tracking, Live-Reporting", dict(size=9.5, space_after=6)),
    ("KI: Now Assist for SPM",
     dict(size=10, bold=True, color=ORANGE, space_after=2)),
    ("•  Risiko-Frühwarnung, Ressourcen-Empfehlungen, Agile Story Generation",
     dict(size=9.5)),
]
add_text(x2 + 0.8, col_top + 4.6, cw - 1.6, col_h - 5, c2_paragraphs)

# --- Column 3: STIHL-Kontext ---
x3 = x2 + cw + gap
add_rect(x3, col_top, cw, 3.8, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(x3, col_top, cw, 3.8, [
    ("ServiceNow @ STIHL",
     dict(size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
], anchor=MSO_ANCHOR.MIDDLE)
add_rect(x3, col_top + 4, cw, col_h - 4, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

c3_paragraphs = [
    ("Etabliert, geprüft, im Betrieb",
     dict(size=10.5, bold=True, color=DARK, space_after=6)),
    ("Im Einsatz seit 2014", dict(size=10, bold=True, color=GREEN, space_after=2)),
    ("•  Globaler IT Service Management-Standard – inklusive China",
     dict(size=9.5, space_after=6)),
    ("DSGVO-konformes Hosting", dict(size=10, bold=True, color=GREEN, space_after=2)),
    ("•  Datacenter Frankfurt am Main", dict(size=9.5)),
    ("•  Datacenter Düsseldorf",
     dict(size=9.5, space_after=2)),
    ("•  Vertragliche und technische DSGVO-Konformität gegeben",
     dict(size=9.5, space_after=6)),
    ("Bestehende STIHL-Basis", dict(size=10, bold=True, color=GREEN, space_after=2)),
    ("•  Etablierte Plattform, Betriebsteam, Lieferantenbeziehung",
     dict(size=9.5)),
    ("•  ITSM-CMDB als Andockpunkt für SPM (Apps, Services, Verträge)",
     dict(size=9.5, space_after=8)),
    ("Konsequenz",
     dict(size=10.5, bold=True, color=ORANGE, space_after=2)),
    ("•  Plattform-Erweiterung statt Tool-Neubeschaffung",
     dict(size=10, italic=True, color=DARK)),
    ("•  Synergien bei Lizenz, Betrieb, Schulung, Governance",
     dict(size=10, italic=True, color=DARK)),
]
add_text(x3 + 0.8, col_top + 4.6, cw - 1.6, col_h - 5, c3_paragraphs)

# ============== REFERENZKUNDEN ==============
ref_top = 48
add_text(4, ref_top, total, 2.4, [
    ("Große deutsche Referenzkunden – ServiceNow im Einsatz in IT und Geschäftsprozessen",
     dict(size=13, bold=True, color=DARK)),
])

# 8 customer placeholder boxes (logos to be inserted manually)
customers = [
    ("Siemens",          "Industrie / Tech"),
    ("SAP",              "Software"),
    ("Allianz",          "Versicherung"),
    ("Deutsche Bank",    "Banking"),
    ("Deutsche Telekom", "Telekommunikation"),
    ("Bosch",            "Industrie / Automotive"),
    ("Mercedes-Benz",    "Automotive"),
    ("Volkswagen",       "Automotive"),
]
box_top = 51
box_h = 8
n = len(customers)
gap_c = 1
box_w = (total - (n - 1) * gap_c) / n

for i, (name, branche) in enumerate(customers):
    x = 4 + i * (box_w + gap_c)
    add_rect(x, box_top, box_w, box_h, WHITE, line=ORANGE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # Logo placeholder area (top half)
    add_text(x, box_top + 0.6, box_w, 3.0, [
        ("[Logo]", dict(size=9, italic=True, color=LIGHT_GREY, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    # name + branche (bottom half)
    add_text(x, box_top + 3.6, box_w, 2.4, [
        (name, dict(size=10.5, bold=True, color=DARK, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x, box_top + 5.8, box_w, 2.0, [
        (branche, dict(size=8.5, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)

# Reference statement bar
add_rect(4, 60.5, total, 4, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(5.5, 60.5, total - 3, 4, [
    ("Mehr als 85 % der Fortune 500 und nahezu alle DAX-Konzerne nutzen ServiceNow – darunter führende deutsche Industrie-, Finanz- und Telekommunikations-Unternehmen.",
     dict(size=11, color=WHITE, italic=True)),
], anchor=MSO_ANCHOR.MIDDLE)

# ============== KEY MESSAGES ==============
km_top = 67
add_text(4, km_top, total, 2.2, [
    ("Kernbotschaften", dict(size=13, bold=True, color=DARK)),
])

bullets = [
    ("Bewährter Industriestandard",
     "Etabliertes Plattform-Ökosystem; Marktführer Gartner Magic Quadrant SPM."),
    ("Pro-Modul = direkter STIHL-Fit",
     "SPM Pro deckt Strategy → Plan → Execute → Deliver durchgängig ab."),
    ("Plattform schon im Haus",
     "Seit 2014 global im Einsatz, DSGVO-konform in DE gehostet – kein Neuaufbau."),
]
bb_top = 69.5
bw = (total - 2 * 1.5) / 3
for i, (head, body) in enumerate(bullets):
    x = 4 + i * (bw + 1.5)
    add_rect(x, bb_top, bw, 7.0, WHITE, line=ORANGE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # number circle
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                               ex(x + 0.7), ex(bb_top + 0.7),
                               ex(2.4), ex(2.4))
    c.fill.solid(); c.fill.fore_color.rgb = ORANGE
    c.line.fill.background()
    set_paragraphs(c, [(str(i + 1), dict(size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER))])
    c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_text(x + 3.6, bb_top + 0.5, bw - 4.0, 2.4, [
        (head, dict(size=11, bold=True, color=DARK)),
    ])
    add_text(x + 3.6, bb_top + 2.6, bw - 4.0, 4.5, [
        (body, dict(size=10, color=GREY_TXT)),
    ])

# Footer
add_text(4, 86.3, total, 2.2, [
    ("Quellen: ServiceNow Produktdokumentation; STIHL IT (Bestand seit 2014, Hosting Frankfurt + Düsseldorf); öffentlich genannte Referenzkunden.",
     dict(size=8.5, italic=True, color=GREY_TXT)),
])
add_rect(0, 89.3, 160, 0.7, ORANGE)

# ============== Move slide right after slide 8 ==============
sldIdLst = prs.slides._sldIdLst
slides = list(sldIdLst)
new_id = slides[-1]  # the just-added slide
# desired position: index 8 (right after slide 8 which is at index 7)
target_index = 8
sldIdLst.remove(new_id)
sldIdLst.insert(target_index, new_id)

prs.save(DST)
print("Saved:", DST)
