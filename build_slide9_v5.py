"""Slide 9 v9: add SPM Pro drawer below the Steckbrief.
Reference customers and STIHL section shift down.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v8.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v9.pptx"

ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)

NOW_BLACK    = RGBColor(0x03, 0x2D, 0x42)
NOW_BLACK_2  = RGBColor(0x05, 0x3B, 0x55)
NOW_GREEN    = RGBColor(0x62, 0xD8, 0x4E)
NOW_GREEN_DK = RGBColor(0x29, 0x9B, 0x32)
NOW_GREY     = RGBColor(0xB6, 0xC0, 0xC8)

U = 76200
def ex(u): return int(u * U)

prs = Presentation(SRC)
slide = prs.slides[8]

for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def add_oval(left, top, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s

def add_triangle(left, top, w, h, fill, shape=MSO_SHAPE.ISOSCELES_TRIANGLE):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK, default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(50000); tf.margin_right = Emu(50000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)
        if "space_after" in st:
            p.space_after = Pt(st["space_after"])

def add_text(left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.TOP, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ============== Slide-level header ==============
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("ServiceNow als Tool – Reichweite, Referenzkunden und STIHL-Kontext",
     dict(size=20, bold=True, color=WHITE)),
])
add_text(3, 3.2, 140, 2.4, [
    ("Bewährter Industriestandard, große deutsche Referenzbasis und seit 2014 bei STIHL etabliert",
     dict(size=12, italic=True, color=LIGHT_GREY)),
])
add_rect(150, 0, 10, 6, ORANGE)
tb = slide.shapes.add_textbox(ex(150), ex(0), ex(10), ex(6))
tb.fill.background(); tb.line.fill.background()
tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_paragraphs(tb, [("Folie 9", dict(size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER))])

# ============== STECKBRIEF (compact) ==============
sb_x, sb_y = 4, 7.0
sb_w, sb_h = 152, 21.0

add_rect(sb_x, sb_y, sb_w, sb_h, NOW_BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_rect(sb_x, sb_y, 0.9, sb_h, NOW_GREEN)

# Hero band
hero_y = sb_y + 0.5
add_oval(sb_x + 2.6, hero_y + 0.6, 1.4, 1.4, NOW_GREEN)
add_text(sb_x + 4.6, hero_y, 60, 2.8, [
    ("servicenow", dict(size=20, bold=True, color=WHITE)),
], anchor=MSO_ANCHOR.MIDDLE)
add_text(sb_x + 38, hero_y, 110, 2.8, [
    ("The AI Platform for Business Transformation",
     dict(size=11, italic=True, color=NOW_GREEN, align=PP_ALIGN.RIGHT)),
], anchor=MSO_ANCHOR.MIDDLE)
add_text(sb_x + 2.6, hero_y + 3.0, sb_w - 4, 2.0, [
    ("ServiceNow ist heute weit mehr als ein Tool für IT-Tickets – die Now Platform verbindet Abteilungen, Daten und Systeme auf einer Cloud-Plattform für digitale Geschäftsabläufe.",
     dict(size=10, italic=True, color=NOW_GREY)),
])

# 3 Prinzipien
prin_y = sb_y + 5.7
prin_h = 2.0
principles = [
    ("Zentralisierung", "Eine Oberfläche statt viele Insel-Apps"),
    ("Automatisierung", "Routine durch KI & Workflows ersetzen"),
    ("User Experience", "Moderne, App-ähnliche Bedienung"),
]
inner_x = sb_x + 2.6
inner_w = sb_w - 5.2
gap_p = 1.0
pw = (inner_w - 2 * gap_p) / 3
for i, (head, sub) in enumerate(principles):
    x = inner_x + i * (pw + gap_p)
    add_rect(x, prin_y, pw, prin_h, NOW_BLACK_2, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line=NOW_GREEN, line_w=0.5)
    add_oval(x + 0.5, prin_y + 0.6, 0.7, 0.7, NOW_GREEN)
    add_text(x + 1.6, prin_y, 14, prin_h, [
        (head, dict(size=10, bold=True, color=WHITE)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x + 15.5, prin_y, pw - 16, prin_h, [
        (sub, dict(size=8.8, italic=True, color=NOW_GREY)),
    ], anchor=MSO_ANCHOR.MIDDLE)

# 4 Workflow-Säulen
pillar_label_y = prin_y + prin_h + 0.7
add_text(inner_x, pillar_label_y, 80, 1.6, [
    ("DAS HAUPTPORTFOLIO – DIE WORKFLOW-SÄULEN",
     dict(size=8.5, bold=True, color=NOW_GREEN)),
])

pillars = [
    ("IT Workflows", "der Ursprung", [
        "ITSM – Ticketing, Incident, Change",
        "ITOM – Infrastruktur-Monitoring",
        "SPM – Strategic Portfolio Mgmt ★",
    ]),
    ("Employee Workflows", "Mitarbeiter", [
        "HR Service Delivery & Onboarding",
        "Workplace Service Delivery",
        "Self-Service & Mobile",
    ]),
    ("Customer Workflows", "Kunden", [
        "Customer Service Management",
        "Field Service Management",
        "Vernetzung Service ↔ Technik",
    ]),
    ("Creator Workflows", "Eigene Lösungen", [
        "App Engine (Low-Code)",
        "Eigene Apps durch Fachbereich",
        "Workflow-Studio & Templates",
    ]),
]
pillar_top = pillar_label_y + 1.8
pillar_h = 8.5
gap_pi = 1.0
piw = (inner_w - 3 * gap_pi) / 4

for i, (head, sub, items) in enumerate(pillars):
    x = inner_x + i * (piw + gap_pi)
    add_rect(x, pillar_top, piw, pillar_h, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(x, pillar_top, piw, 2.8, NOW_GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_oval(x + 0.6, pillar_top + 0.6, 1.6, 1.6, NOW_BLACK)
    add_text(x + 0.6, pillar_top + 0.6, 1.6, 1.6, [
        (str(i + 1), dict(size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x + 2.6, pillar_top + 0.2, piw - 3, 1.4, [
        (head, dict(size=10, bold=True, color=NOW_BLACK)),
    ], anchor=MSO_ANCHOR.TOP)
    add_text(x + 2.6, pillar_top + 1.4, piw - 3, 1.3, [
        (sub, dict(size=8.2, italic=True, color=NOW_BLACK)),
    ], anchor=MSO_ANCHOR.TOP)
    for j, item in enumerate(items):
        iy = pillar_top + 3.2 + j * 1.65
        is_spm = "SPM" in item
        add_oval(x + 0.7, iy + 0.5, 0.5, 0.5, NOW_GREEN_DK if not is_spm else ORANGE)
        add_text(x + 1.5, iy, piw - 2, 1.5, [
            (item.replace(" ★",""),
             dict(size=8.8, bold=is_spm,
                  color=NOW_BLACK if not is_spm else ORANGE)),
        ], anchor=MSO_ANCHOR.MIDDLE)

# AI banner
ai_top = pillar_top + pillar_h + 0.6
ai_h = 2.2
add_rect(inner_x, ai_top, inner_w, ai_h, NOW_GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
badge_w = 7
add_rect(inner_x + 0.4, ai_top + 0.25, badge_w, ai_h - 0.5, NOW_BLACK,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(inner_x + 0.4, ai_top + 0.25, badge_w, ai_h - 0.5, [
    ("KI 2024 / 25", dict(size=8.5, bold=True, color=NOW_GREEN, align=PP_ALIGN.CENTER)),
], anchor=MSO_ANCHOR.MIDDLE)
add_text(inner_x + badge_w + 1.2, ai_top, inner_w - badge_w - 1.6, ai_h, [
    ("Aktueller Fokus: Generative AI (Zusammenfassungen, Code, Vorschläge) und Agentic AI (autonome KI-Agenten erledigen Aufgaben selbständig).",
     dict(size=9.5, bold=True, color=NOW_BLACK)),
], anchor=MSO_ANCHOR.MIDDLE)

# ============== SPM PRO DRAWER ==============
# Connector triangle from Steckbrief down to drawer (visual cue: drawer pulled out)
drawer_top = sb_y + sb_h + 0.5  # 28.5
# small green connector tab
tab_w = 30
tab_x = inner_x  # under IT pillar
add_triangle(tab_x + tab_w/2 - 1.0, drawer_top - 1.4, 2.0, 1.4, NOW_GREEN)
add_rect(tab_x, drawer_top - 0.2, tab_w, 0.4, NOW_GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

drawer_h = 11.5
add_rect(sb_x, drawer_top, sb_w, drawer_h, NOW_BLACK_2, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line=NOW_GREEN, line_w=1.2)
# Drawer title bar
add_text(sb_x + 3, drawer_top + 0.5, 80, 2.4, [
    ("DETAIL  ▾  ServiceNow SPM Pro – das STIHL-relevante Modul",
     dict(size=11, bold=True, color=NOW_GREEN)),
])
add_text(sb_x + 80, drawer_top + 0.5, sb_w - 84, 2.4, [
    ("SPM Pro = SPM Standard  +  Strategic Planning  +  Agile/EAP  +  Pro-Plattform-Features",
     dict(size=9.5, italic=True, color=NOW_GREY, align=PP_ALIGN.RIGHT)),
])

# 6 capability tiles
tiles = [
    ("Strategic Planning",    "OKR-Framework, Goals, Roadmaps, Investment Funding",      True),
    ("Demand & Portfolio",    "Erfassen, bewerten, priorisieren – ein Workspace",        False),
    ("Project Management",    "Klassisch (Gantt) und agil (Kanban) in einem System",     False),
    ("Resource & Financial",  "Skill-Planung, Forecasts, Scenario / What-if-Planung",    True),
    ("Agile / Hybrid (EAP)",  "Enterprise Agile Planning, hybride Roadmaps, Boards",     True),
    ("Now Assist for SPM",    "Risiko-Frühwarnung, Story Generation, KI-Reporting",      True),
]
tile_top = drawer_top + 3.2
tile_h = 7.5
gap_t = 1.0
inner_d_w = sb_w - 6
tw = (inner_d_w - 5 * gap_t) / 6

for i, (head, body, pro_only) in enumerate(tiles):
    x = sb_x + 3 + i * (tw + gap_t)
    add_rect(x, tile_top, tw, tile_h, NOW_BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line=NOW_GREEN, line_w=0.5)
    add_rect(x, tile_top, tw, 0.35, NOW_GREEN)
    # Pro badge
    if pro_only:
        add_rect(x + tw - 4.2, tile_top + 0.7, 3.6, 1.4, NOW_GREEN,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(x + tw - 4.2, tile_top + 0.7, 3.6, 1.4, [
            ("PRO", dict(size=7.5, bold=True, color=NOW_BLACK, align=PP_ALIGN.CENTER)),
        ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x + 0.5, tile_top + 0.8, tw - 5.2 if pro_only else tw - 1, 2.6, [
        (head, dict(size=10, bold=True, color=WHITE)),
    ], anchor=MSO_ANCHOR.TOP)
    add_text(x + 0.5, tile_top + 3.2, tw - 1, tile_h - 3.4, [
        (body, dict(size=8.7, color=NOW_GREY)),
    ], anchor=MSO_ANCHOR.TOP)

# ============== MIDDLE: Reference customers (shifted down) ==============
ref_top = drawer_top + drawer_h + 1.5  # ~41.5
total = 152
add_text(4, ref_top, total, 2.4, [
    ("Große deutsche Referenzkunden – ServiceNow im Einsatz in IT und Geschäftsprozessen",
     dict(size=12.5, bold=True, color=DARK)),
])

customers = [
    ("SIEMENS",          "Industrie / Tech",       RGBColor(0x00, 0x9D, 0x9D), WHITE),
    ("SAP",              "Software",               RGBColor(0x00, 0x3A, 0x5D), RGBColor(0x07, 0xC8, 0xF0)),
    ("Allianz",          "Versicherung",           RGBColor(0x00, 0x37, 0x81), WHITE),
    ("Deutsche Bank",    "Banking",                RGBColor(0x00, 0x18, 0xA8), WHITE),
    ("T",                "Deutsche Telekom",       RGBColor(0xE2, 0x00, 0x74), WHITE),
    ("BOSCH",            "Industrie / Automotive", RGBColor(0xC8, 0x10, 0x2E), WHITE),
    ("Mercedes-Benz",    "Automotive",             RGBColor(0x00, 0x00, 0x00), WHITE),
    ("Volkswagen",       "Automotive",             RGBColor(0x00, 0x1E, 0x50), WHITE),
]
box_top = ref_top + 2.6
box_h_logo = 6.0
box_h_total = 9.5
n = len(customers)
gap_c = 0.8
box_w = (total - (n - 1) * gap_c) / n

for i, (name, branche, bg, fg) in enumerate(customers):
    x = 4 + i * (box_w + gap_c)
    add_rect(x, box_top, box_w, box_h_logo, bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    name_size = 14 if len(name) <= 4 else (12 if len(name) <= 8 else 10.5)
    add_text(x, box_top, box_w, box_h_logo, [
        (name, dict(size=name_size, bold=True, color=fg, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x, box_top + box_h_logo + 0.2, box_w, 3.0, [
        (branche, dict(size=8.2, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.TOP)

div_y = box_top + box_h_total + 1.5
add_text(4, div_y, total, 1.6, [
    ("Stilisierte Markenfarben als Logo-Platzhalter – echte Logos können in der finalen Vorstandsfassung eingesetzt werden  |  Quelle: ServiceNow Customer References",
     dict(size=7.8, italic=True, color=LIGHT_GREY)),
])
add_rect(4, div_y + 1.8, total, 0.15, ORANGE)

# ============== BOTTOM: STIHL ==============
bot_top = div_y + 2.5
add_text(4, bot_top, total, 2.0, [
    ("ServiceNow @ STIHL – die Plattform ist bereits etabliert",
     dict(size=13, bold=True, color=DARK)),
])

stihl_boxes = [
    ("Im Einsatz seit 2014",
     "Globaler IT Service Management-Standard – inklusive China.\n"
     "Etabliertes Betriebsteam, eingespielte Lieferantenbeziehung.",
     "12 Jahre Betrieb"),
    ("DSGVO-konformes Hosting in Deutschland",
     "Hosting in den ServiceNow-Datacentern Frankfurt am Main und Düsseldorf.\n"
     "Vertragliche und technische DSGVO-Konformität gegeben.",
     "Frankfurt + Düsseldorf"),
    ("Anschluss-/Erweiterungsfähigkeit",
     "ITSM-CMDB als Andockpunkt für SPM (Apps, Services, Verträge).\n"
     "Plattform-Erweiterung statt Tool-Neubeschaffung – Synergien bei Lizenz/Betrieb.",
     "Plattform-Synergie"),
]
sb_top2 = bot_top + 2.4
# Make STIHL boxes more compact
sb_h2 = 86.0 - sb_top2 - 4  # leave room for footer
gap_s = 2
sw2 = (total - 2 * gap_s) / 3
for i, (head, body, badge) in enumerate(stihl_boxes):
    x = 4 + i * (sw2 + gap_s)
    add_rect(x, sb_top2, sw2, sb_h2, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(x, sb_top2, sw2, 0.8, GREEN)
    badge_w = 15
    add_rect(x + sw2 - badge_w - 0.6, sb_top2 + 1.2, badge_w, 2.2, GREEN,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x + sw2 - badge_w - 0.6, sb_top2 + 1.2, badge_w, 2.2, [
        (badge, dict(size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x + 1.2, sb_top2 + 1.4, sw2 - badge_w - 2.4, 4.0, [
        (head, dict(size=11, bold=True, color=DARK)),
    ])
    body_paras = [(line, dict(size=9.5, color=DARK, space_after=3))
                  for line in body.split("\n") if line]
    add_text(x + 1.2, sb_top2 + 5.6, sw2 - 2.4, sb_h2 - 6.4, body_paras)

# FOOTER
add_text(4, 86.5, total, 2.0, [
    ("Quellen: ServiceNow (servicenow.com), Customer References; STIHL IT (Bestand seit 2014, Hosting Frankfurt + Düsseldorf, DSGVO-konform).",
     dict(size=8, italic=True, color=GREY_TXT)),
])
add_rect(0, 89.3, 160, 0.7, ORANGE)

prs.save(DST)
print("Saved:", DST)
