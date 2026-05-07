"""Replace slide 8 of the SNOW SPM board deck with a native, board-grade
PowerPoint layout describing the ServiceNow SPM scope."""
from copy import deepcopy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC = "/root/.claude/uploads/07f3ab55-e569-4094-81eb-8406d6bf9995/dc356137-2026_05_18_VS_Zielbild_SNOW_SPM.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v2.pptx"

# Colors (matching the mockup)
ORANGE      = RGBColor(0xF0, 0x7F, 0x12)
ORANGE2     = RGBColor(0xE9, 0x6C, 0x0C)
ORANGE3     = RGBColor(0xC9, 0x53, 0x00)
ORANGE4     = RGBColor(0x9C, 0x3F, 0x08)
PHASE = [ORANGE, ORANGE2, ORANGE3, ORANGE4]
DARK        = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT    = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG    = RGBColor(0xF7, 0xF7, 0xF7)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
FOUND       = RGBColor(0x2C, 0x3E, 0x50)
RED         = RGBColor(0xB9, 0x1C, 0x1C)
LIGHT_GREY  = RGBColor(0xD1, 0xD5, 0xDB)

# Slide is 13.333" x 7.5"  → 12192000 x 6858000 EMU
SW = 12192000
SH = 6858000
# Mockup grid was 160 x 90  → unit = 76200 EMU
U = 76200

def emu_x(u): return int(u * U)
def emu_y(u): return int(u * U)

prs = Presentation(SRC)
slide = prs.slides[7]  # slide 8 (0-indexed)

# --- Clear all existing shapes on the slide ---
spTree = slide.shapes._spTree
for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left_u, top_u, w_u, h_u, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, emu_x(left_u), emu_y(top_u), emu_x(w_u), emu_y(h_u))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.2)
    s.shadow.inherit = False
    return s

def set_text(shape, text, *, size=11, bold=False, italic=False, color=DARK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, font="Calibri"):
    tf = shape.text_frame
    tf.margin_left = Emu(50000)
    tf.margin_right = Emu(50000)
    tf.margin_top = Emu(20000)
    tf.margin_bottom = Emu(20000)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(text, str):
        runs = [(text, dict(size=size, bold=bold, italic=italic, color=color, font=font))]
    else:
        runs = text  # list of (text, style-dict)
    first = True
    for txt, st in runs:
        r = p.add_run() if not first else p.add_run()
        r.text = txt
        r.font.name = st.get("font", font)
        r.font.size = Pt(st.get("size", size))
        r.font.bold = st.get("bold", bold)
        r.font.italic = st.get("italic", italic)
        r.font.color.rgb = st.get("color", color)
        first = False

def add_text(left_u, top_u, w_u, h_u, text, **kw):
    tb = slide.shapes.add_textbox(emu_x(left_u), emu_y(top_u), emu_x(w_u), emu_y(h_u))
    tb.fill.background()
    tb.line.fill.background()
    set_text(tb, text, **kw)
    return tb

# ============== HEADER ==============
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6,
         "ServiceNow SPM – Eine Plattform von der Strategie bis zur Umsetzung",
         size=22, bold=True, color=WHITE, anchor=MSO_ANCHOR.TOP)
add_text(3, 3.2, 140, 2.4,
         "Integrierter Leistungsumfang zur Ablösung fragmentierter Einzeltools",
         size=13, italic=True, color=LIGHT_GREY, anchor=MSO_ANCHOR.TOP)
add_rect(150, 0, 10, 6, ORANGE)
add_text(150, 0, 10, 6, "Folie 8", size=12, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ============== PHASE BAR ==============
phases = [
    ("STRATEGY",      "Ziele setzen"),
    ("PLAN",          "Priorisieren"),
    ("EXECUTE",       "Umsetzen"),
    ("DELIVER & RUN", "Wertbeitrag sichern"),
]
x_left, x_right = 4, 156
total_w = x_right - x_left
phase_w = total_w / 4
phase_y = 8.5
phase_h = 5.5

for i, (title, sub) in enumerate(phases):
    x = x_left + i * phase_w
    add_rect(x + 0.4, phase_y, phase_w - 0.8, phase_h, PHASE[i],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x, phase_y + 0.4, phase_w, 2.4, title,
             size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    add_text(x, phase_y + 2.8, phase_w, 2.4, sub,
             size=11, italic=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    if i < 3:
        ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                    emu_x(x + phase_w - 0.6), emu_y(phase_y + phase_h/2 - 0.7),
                                    emu_x(1.2), emu_y(1.4))
        ar.fill.solid(); ar.fill.fore_color.rgb = DARK
        ar.line.fill.background()

# ============== MODULE TILES ==============
modules = [
    ["Strategic Planning",     "Goal Framework / OKR",    "Investment Funding"],
    ["Demand Management ★",    "Portfolio Management",    "Innovation / Idea Mgmt"],
    ["Project Management",     "Resource Management",     "Agile / SAFe Workflows"],
    ["Application Portfolio",  "Benefits / Outcomes",     "Reporting & Dashboards"],
]
tile_top = 14.6
tile_h = 5.0
tile_gap = 1.0

for i, mods in enumerate(modules):
    x = x_left + i * phase_w
    for j, m in enumerate(mods):
        ty = tile_top + j * (tile_h + tile_gap)
        s = add_rect(x + 1.0, ty, phase_w - 2.0, tile_h, WHITE, line=PHASE[i],
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        is_connector = "★" in m
        add_text(x + 1.0, ty, phase_w - 2.0, tile_h, m,
                 size=12, bold=is_connector,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 color=DARK)

# Connector legend
add_text(x_left, 33.4, total_w, 2.0,
         "★ Demand = Connector zwischen Backlog und Portfolio  (vgl. Zielbild PPM, Folie 22)",
         size=10, italic=True, color=ORANGE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ============== FOUNDATION BAR ==============
fy = 36
fh = 5.5
add_rect(x_left, fy, total_w, fh, FOUND, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(x_left + 2.5, fy + 0.5, 40, 2.4, "Plattform-Fundament",
         size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.TOP)
add_text(x_left + 2.5, fy + 2.7, total_w - 5, 2.4,
         "Eine Datenbasis  ·  Workflows  ·  Rollenmodell  ·  Dashboards  ·  KI / Now Assist  ·  Integrationen (SAP, PIT, EV)",
         size=11, color=RGBColor(0xE5, 0xE7, 0xEB), anchor=MSO_ANCHOR.TOP)

# ============== PAIN-POINT 2x2 ==============
pp_top = 43
add_text(x_left, pp_top, total_w, 2.4,
         "Was leistet ServiceNow SPM – und welchen STIHL-Pain-Point löst es?",
         size=14, bold=True, color=DARK, anchor=MSO_ANCHOR.TOP)

rows = [
    ("Strategy",      "Strategische Ziele, OKRs und Investitionsrahmen direkt mit dem Portfolio verknüpft.",
                      "heute: keine durchgängige Priorisierungslogik"),
    ("Plan",          "Demands, Ideen und Projekte zentral erfassen, bewerten und priorisieren.",
                      "heute: Parallel-Prozesse, intransparente Gremien"),
    ("Execute",       "Klassisches, agiles und hybrides PM inkl. Ressourcen- und Finanzsteuerung in einem System.",
                      "heute: PIT / MSPO abgekündigt, manuelle Datenhaltung"),
    ("Deliver & Run", "Nutzen-Tracking, Anwendungsportfolio und Live-Reporting für Vorstand und Bereiche.",
                      "heute: bereichsübergreifende Intransparenz"),
]
grid_top = 45.6
cell_w = total_w / 2 - 1
cell_h = 7.4
for idx, (phase, leistung, pain) in enumerate(rows):
    col = idx % 2
    row = idx // 2
    cx = x_left + col * (cell_w + 2)
    cy = grid_top + row * (cell_h + 1.0)
    add_rect(cx, cy, cell_w, cell_h, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(cx, cy, 0.9, cell_h, PHASE[idx])
    add_text(cx + 1.6, cy + 0.5, cell_w - 2, 2.0, phase,
             size=12, bold=True, color=DARK, anchor=MSO_ANCHOR.TOP)
    add_text(cx + 1.6, cy + 2.4, cell_w - 2, 3.4, leistung,
             size=10.5, color=DARK, anchor=MSO_ANCHOR.TOP)
    add_text(cx + 1.6, cy + cell_h - 2.0, cell_w - 2, 1.8,
             "→ " + pain, size=10, italic=True, color=RED, anchor=MSO_ANCHOR.TOP)

# ============== KEY MESSAGES ==============
km_top = 64
add_text(x_left, km_top, total_w, 2.2, "Kernbotschaften",
         size=14, bold=True, color=DARK, anchor=MSO_ANCHOR.TOP)

bullets = [
    ("End-to-end statt Insellösung",
     "Durchgängiger Prozess von Strategie bis Betrieb – ohne Medienbrüche."),
    ("Standard statt Eigenbau",
     "Releasefähige Module, Best-Practice-Workflows; Phase 1: Demand, Strategy, MDM, Portfolio."),
    ("Plattform statt Tool",
     "Gemeinsame Datenbasis mit ITSM, EV, Workflows – ausbaubar auf weitere Geschäftsthemen."),
]
box_top = 66.2
box_h = 8.6
gap = 1.5
bw = (total_w - 2 * gap) / 3
for i, (head, body) in enumerate(bullets):
    x = x_left + i * (bw + gap)
    add_rect(x, box_top, bw, box_h, WHITE, line=ORANGE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # number circle
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                               emu_x(x + 0.8), emu_y(box_top + 0.8),
                               emu_x(3.0), emu_y(3.0))
    c.fill.solid(); c.fill.fore_color.rgb = ORANGE
    c.line.fill.background()
    set_text(c, str(i + 1), size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(x + 4.2, box_top + 0.7, bw - 4.5, 2.6, head,
             size=12, bold=True, color=DARK, anchor=MSO_ANCHOR.TOP)
    add_text(x + 4.2, box_top + 3.0, bw - 4.5, 5.0, body,
             size=10.5, color=GREY_TXT, anchor=MSO_ANCHOR.TOP)

# ============== FOOTER ==============
add_text(x_left, 86.3, total_w, 2.2,
         "Quelle: ServiceNow SPM Pro/Enterprise – Modul-Detail siehe Backup   |   Gartner Magic Quadrant SPM: Leader 2024 / 2025",
         size=9, italic=True, color=GREY_TXT, anchor=MSO_ANCHOR.TOP)
add_rect(0, 89.3, 160, 0.7, ORANGE)

prs.save(DST)
print("Saved:", DST)
