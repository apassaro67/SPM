"""Vorstands-Vorlage: Governance-Bewertung im Rahmen der Plattformeinführung.

Zeigt, wie der Aspekt 'Governance-Strukturen in den Fachbereichen bewerten, um
Komplexitaet zu vermeiden; Prozess- und FTE-Aufwand im Folgeprojekt bewerten'
erarbeitet und dem Vorstand aufgezeigt wird. Stil an die SPM-Decks angelehnt.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DST = "2026_07_08_Governance_Bewertung_Plattform_Vorstand.pptx"

ORANGE    = RGBColor(0xF0, 0x7F, 0x12); ORANGE_LT = RGBColor(0xFD, 0xE6, 0xCC)
DARK      = RGBColor(0x1F, 0x29, 0x37); GREY_TXT  = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG  = RGBColor(0xF7, 0xF7, 0xF7); ROW_ALT   = RGBColor(0xEE, 0xF1, 0xF4)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF); NAVY      = RGBColor(0x2C, 0x3E, 0x50)
BLUE      = RGBColor(0x2E, 0x5A, 0x8F); BLUE_LT   = RGBColor(0xE4, 0xEC, 0xF4)
PURPLE    = RGBColor(0x6B, 0x4F, 0xA0)
GREEN     = RGBColor(0x2E, 0x7D, 0x32); GREEN_LT  = RGBColor(0xE3, 0xF2, 0xEA)
AMBER     = RGBColor(0xB4, 0x76, 0x09); AMBER_LT  = RGBColor(0xFB, 0xEF, 0xD6)
RED       = RGBColor(0xB9, 0x1C, 0x1C); RED_LT    = RGBColor(0xFB, 0xE4, 0xE1)
LINE_GREY = RGBColor(0xD1, 0xD5, 0xDB)

U = 76200
def ex(u): return int(u * U)
FONT = "Calibri"

prs = Presentation(); prs.slide_width = ex(160); prs.slide_height = ex(90)
BLANK = prs.slide_layouts[6]

def add_rect(slide, left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0, round_=False):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else shape
    s = slide.shapes.add_shape(shp, ex(left), ex(top), ex(w), ex(h))
    if round_:
        try: s.adjustments[0] = 0.08
        except Exception: pass
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=11, default_color=DARK, default_align=PP_ALIGN.LEFT, line_spacing=None):
    tf = shape.text_frame; tf.word_wrap = True
    tf.margin_left = Emu(45000); tf.margin_right = Emu(45000); tf.margin_top = Emu(15000); tf.margin_bottom = Emu(15000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        text, st = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        if line_spacing: p.line_spacing = line_spacing
        if "space_after" in st: p.space_after = Pt(st["space_after"])
        runs = st.get("runs")
        if runs:
            for rt, rs in runs:
                run = p.add_run(); run.text = rt; run.font.name = FONT
                run.font.size = Pt(rs.get("size", st.get("size", default_size)))
                run.font.bold = rs.get("bold", False); run.font.italic = rs.get("italic", False)
                run.font.color.rgb = rs.get("color", st.get("color", default_color))
        else:
            run = p.add_run(); run.text = text; run.font.name = FONT
            run.font.size = Pt(st.get("size", default_size)); run.font.bold = st.get("bold", False)
            run.font.italic = st.get("italic", False); run.font.color.rgb = st.get("color", default_color)

def add_text(slide, left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.TOP, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background(); tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw); return tb

def header(slide, kicker, title, page):
    add_rect(slide, 0, 0, 160, 90, WHITE)
    add_rect(slide, 0, 0, 1.6, 90, ORANGE)
    add_text(slide, 7, 5, 130, 4, [(kicker, {"size": 11, "bold": True, "color": ORANGE})])
    add_text(slide, 7, 8.5, 146, 9, [(title, {"size": 22, "bold": True, "color": DARK})])
    add_rect(slide, 7, 17.5, 146, 0.12, LINE_GREY)
    add_text(slide, 150, 84.5, 8, 4, [(str(page), {"size": 10, "color": LINE_GREY, "align": PP_ALIGN.RIGHT})])
    add_text(slide, 7, 84.5, 110, 4, [("Governance-Bewertung Plattformeinführung · Vorlage Vorstand", {"size": 9, "color": LINE_GREY})])

# ==========================================================================
# SLIDE 1 -- Titel / Kernbotschaft
# ==========================================================================
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, 160, 90, DARK); add_rect(s, 0, 0, 160, 2.4, ORANGE)
add_text(s, 12, 20, 136, 5, [("PLATTFORMEINFÜHRUNG · GOVERNANCE-BEWERTUNG · VORLAGE VORSTAND", {"size": 12, "bold": True, "color": ORANGE})])
add_text(s, 12, 26, 138, 20,
         [("Die Plattform ist nur so gut", {"size": 36, "bold": True, "color": WHITE}),
          ("wie die Governance darunter.", {"size": 36, "bold": True, "color": ORANGE})])
add_text(s, 12, 50, 132, 10,
         [("Neben der Software müssen die Governance-Strukturen in den Fachbereichen bewertet werden — "
           "sonst digitalisieren wir bestehende Komplexität und zahlen sie dauerhaft in Prozess und FTE. "
           "Der resultierende Bedarf wird im Folgeprojekt bewertet.",
           {"size": 14, "color": RGBColor(0xC7, 0xCD, 0xD4)})], line_spacing=1.25)
add_rect(s, 12, 66, 136, 12, RGBColor(0x27, 0x31, 0x40), round_=True)
add_rect(s, 12, 66, 0.9, 12, ORANGE)
add_text(s, 15, 66, 130, 12,
         [("", {"runs": [("Leitprinzip:  ", {"bold": True, "color": ORANGE, "size": 14}),
             ("„Harmonize before you digitize.“  ", {"bold": True, "italic": True, "color": WHITE, "size": 14}),
             ("Standard als Default, lokale Varianten nur mit Begründung (comply-or-explain).",
              {"color": RGBColor(0xC7, 0xCD, 0xD4), "size": 13})]})], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)

# ==========================================================================
# SLIDE 2 -- Der Aspekt & das Risiko (warum Vorstand)
# ==========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "WARUM DER VORSTAND", "Ohne Governance-Bewertung digitalisieren wir Komplexität", 2)
add_text(s, 7, 19.5, 146, 5,
         [("Eine Plattform (z. B. ServiceNow SPM) setzt ein Prozess- und Rollenmodell voraus. Laufen in den "
           "Fachbereichen heute N unterschiedliche Governance-Modelle, entstehen zwei Wege:",
           {"size": 12, "color": GREY_TXT})], line_spacing=1.15)

# zwei Pfade
add_rect(s, 7, 27, 71.5, 24, WHITE, line=RED, line_w=1.25, round_=True)
add_rect(s, 7, 27, 71.5, 5, RED_LT)
add_text(s, 8.8, 27, 68, 5, [("✗  Governance NICHT bewerten", {"size": 12, "bold": True, "color": RED})], anchor=MSO_ANCHOR.MIDDLE)
add_text(s, 9, 33, 68, 17, [(x, {"size": 10, "color": DARK, "space_after": 4}) for x in [
    "•  N Fachbereichs-Modelle 1:1 im Tool abgebildet",
    "•  Dauerhafte Konfigurations- & Betriebskomplexität",
    "•  Höherer, versteckter FTE-Bedarf (je Variante)",
    "•  Schlechte Datenqualität, keine Vergleichbarkeit",
    "•  „Automatisiertes Chaos“ statt Steuerbarkeit"]], line_spacing=1.1)

add_rect(s, 81.5, 27, 71.5, 24, WHITE, line=GREEN, line_w=1.25, round_=True)
add_rect(s, 81.5, 27, 71.5, 5, GREEN_LT)
add_text(s, 83.3, 27, 68, 5, [("✓  Governance bewerten & harmonisieren", {"size": 12, "bold": True, "color": GREEN})], anchor=MSO_ANCHOR.MIDDLE)
add_text(s, 83.5, 33, 68, 17, [(x, {"size": 10, "color": DARK, "space_after": 4}) for x in [
    "•  Ein Standardprozess, wenige begründete Varianten",
    "•  Geringere Betriebskomplexität & -kosten",
    "•  Transparenter, planbarer FTE-Bedarf",
    "•  Konsistente Daten → verlässliche Steuerung",
    "•  Plattform hebt echten Nutzen statt Altlasten"]], line_spacing=1.1)

add_rect(s, 7, 54, 146, 10, ORANGE_LT, round_=True); add_rect(s, 7, 54, 1.0, 10, ORANGE)
add_text(s, 9.5, 54.8, 142, 8,
         [("", {"runs": [("Die Vorstands-Frage lautet nicht „welches Tool“, sondern:  ", {"color": DARK, "size": 12}),
             ("Wollen wir die heutige Governance-Vielfalt zementieren — oder sie im Zuge der Plattform "
              "harmonisieren?", {"bold": True, "color": DARK, "size": 12}),
             ("  Diese Entscheidung determiniert Komplexität, Prozess und FTE für Jahre.",
              {"color": GREY_TXT, "size": 12})]})], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)

add_text(s, 7, 66.5, 146, 3.5, [("Betroffene Governance-Bereiche (Scope der Bewertung)", {"size": 11, "bold": True, "color": DARK})])
scope = ["Demand / Intake", "Portfolio-Priorisierung", "Projekt-/Stage-Gate-Governance",
         "Ressourcen- & Kapazitätsmgmt", "Reporting & KPIs", "Rollen & Gremien (RACI)"]
cw = 23.6
for i, sc in enumerate(scope):
    lx = 7 + (i % 3) * (cw * 2 + 1.5) if False else 7 + (i % 3) * (48.3)
    ly = 70.5 + (i // 3) * 6.5
    add_rect(s, 7 + (i % 3) * 49, ly, 47, 5.5, BLUE_LT, round_=True)
    add_text(s, 7 + (i % 3) * 49, ly, 47, 5.5, [(sc, {"size": 10, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER})], anchor=MSO_ANCHOR.MIDDLE)

# ==========================================================================
# SLIDE 3 -- Wie wird es erarbeitet (Vorgehen)
# ==========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "WIE ERARBEITEN", "Fünf Schritte — vom Ist der Fachbereiche zum FTE-Bedarf", 3)
steps = [
    ("1", "Governance-Ist je Fachbereich", "Wie steuert jeder Fachbereich heute Demand, Portfolio, Projekte, Ressourcen, Reporting? Rollen, Gremien, Tools, Reifegrad — per Assessment-Raster."),
    ("2", "Plattform-Standard (Soll)", "Welches Prozess-, Rollen- und Datenmodell setzt die Plattform voraus bzw. ermöglicht sie? Referenz-Governance definieren."),
    ("3", "Gap- & Komplexitätsanalyse", "Delta Ist ↔ Soll je Fachbereich. Wo Standard möglich (harmonisieren) vs. wo begründete lokale Varianz? Komplexitäts-Heatmap."),
    ("4", "Prozess- & Rollen-Zielbild", "Standardprozesse + Governance-Rollen (Demand-, Portfolio-Mgr, PMO, Ressourcen-Mgmt) ableiten; Harmonisierungs-Entscheidungen treffen."),
    ("5", "Aufwands- & FTE-Bewertung", "FTE für Betrieb der Governance-Rollen + einmaliger Change-Aufwand; Szenarien harmonisiert ↔ föderiert ↔ Status quo."),
]
sw = 28.5; sgap = 1.4
for i, (n, t, d) in enumerate(steps):
    lx = 7 + i * (sw + sgap)
    add_rect(s, lx, 24, sw, 44, WHITE, line=LINE_GREY, line_w=0.9, round_=True)
    add_rect(s, lx, 24, sw, 8, BLUE if i < 4 else ORANGE, round_=True)
    add_rect(s, lx, 27, sw, 5, BLUE if i < 4 else ORANGE)
    add_text(s, lx, 24, sw, 8, [(f"Schritt {n}", {"size": 10.5, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER})], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, lx + 1.4, 33, sw - 2.8, 6, [(t, {"size": 10, "bold": True, "color": DARK})], line_spacing=1.0)
    add_text(s, lx + 1.4, 40, sw - 2.8, 27, [(d, {"size": 8.6, "color": GREY_TXT})], line_spacing=1.1)
    if i < 4:
        add_text(s, lx + sw - 0.5, 44, 2.5, 4, [("→", {"size": 16, "bold": True, "color": ORANGE, "align": PP_ALIGN.CENTER})])

add_rect(s, 7, 71, 146, 12, LIGHT_BG, round_=True)
add_text(s, 9, 72, 142, 3.5, [("Bewertungsdimensionen je Fachbereich (Rating: Reifegrad · Harmonisierungspotenzial · Komplexitätsrisiko)", {"size": 10, "bold": True, "color": DARK})])
add_text(s, 9, 75.8, 142, 6,
         [("Demand/Intake  ·  Portfolio-Priorisierung  ·  Projekt-/Stage-Gate  ·  Ressourcen/Kapazität  ·  "
           "Reporting/KPI & Datenqualität  ·  Rollen & Gremien (RACI)  ·  Standardisierungs-/Toolnutzungsgrad",
           {"size": 9.5, "color": GREY_TXT})], line_spacing=1.2)

# ==========================================================================
# SLIDE 4 -- Ergebnisform: Komplexitäts-Heatmap
# ==========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "ERGEBNISFORM", "So sieht das Ergebnis aus: eine Komplexitäts-Heatmap", 4)
add_text(s, 7, 19.5, 146, 4,
         [("Je Fachbereich × Governance-Dimension: harmonisierbar, begründete Varianz oder Klärungsbedarf. "
           "Macht Komplexitätstreiber und Harmonisierungspotenzial sichtbar. (Darstellung illustrativ — im Folgeprojekt zu erheben.)",
           {"size": 10.5, "color": GREY_TXT})], line_spacing=1.1)

dims = ["Demand", "Portfolio", "Stage-Gate", "Ressourcen", "Reporting", "Rollen/RACI"]
fbs = ["Fachbereich A", "Fachbereich B", "Fachbereich C", "Fachbereich D", "Fachbereich E", "Fachbereich F"]
# 0=green(Standard), 1=amber(Varianz), 2=red(Lücke)
grid = [
    [0,0,1,0,1,0],
    [0,1,1,2,1,1],
    [0,0,0,1,0,0],
    [1,2,2,2,2,1],
    [0,1,0,1,1,0],
    [1,1,2,1,2,2],
]
cmap = {0: (GREEN_LT, GREEN, "S"), 1: (AMBER_LT, AMBER, "V"), 2: (RED_LT, RED, "L")}
gx, gy = 7, 26; lblw = 26; cw = (146 - lblw) / len(dims); hh = 6
add_rect(s, gx, gy, lblw, hh, DARK); add_text(s, gx + 0.6, gy, lblw - 1, hh, [("Fachbereich", {"size": 9, "bold": True, "color": WHITE})], anchor=MSO_ANCHOR.MIDDLE)
for j, d in enumerate(dims):
    cxx = gx + lblw + j * cw
    add_rect(s, cxx, gy, cw, hh, DARK)
    add_text(s, cxx, gy, cw, hh, [(d, {"size": 8.5, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER})], anchor=MSO_ANCHOR.MIDDLE)
rh = 6.2
for i, fb in enumerate(fbs):
    ry = gy + hh + i * rh
    add_rect(s, gx, ry, lblw, rh, WHITE if i % 2 == 0 else ROW_ALT)
    add_text(s, gx + 0.8, ry, lblw - 1, rh, [(fb, {"size": 9.3, "bold": True, "color": DARK})], anchor=MSO_ANCHOR.MIDDLE)
    for j, v in enumerate(grid[i]):
        cxx = gx + lblw + j * cw
        bg, fg, lab = cmap[v]
        add_rect(s, cxx + 0.4, ry + 0.4, cw - 0.8, rh - 0.8, bg, round_=True)
        add_text(s, cxx, ry, cw, rh, [(lab, {"size": 9, "bold": True, "color": fg, "align": PP_ALIGN.CENTER})], anchor=MSO_ANCHOR.MIDDLE)
# Legende
ly = gy + hh + len(fbs) * rh + 2
legend = [(GREEN_LT, GREEN, "S", "Standard-fähig (harmonisieren)"),
          (AMBER_LT, AMBER, "V", "begründete lokale Varianz"),
          (RED_LT, RED, "L", "Lücke / Klärungsbedarf")]
lx = 7
for bg, fg, lab, txt in legend:
    add_rect(s, lx, ly, 4, 4, bg, round_=True); add_text(s, lx, ly, 4, 4, [(lab, {"size": 9, "bold": True, "color": fg, "align": PP_ALIGN.CENTER})], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, lx + 4.7, ly - 0.4, 42, 4.6, [(txt, {"size": 9.5, "color": GREY_TXT})], anchor=MSO_ANCHOR.MIDDLE)
    lx += 49
add_rect(s, 7, ly + 6.5, 146, 8.5, ORANGE_LT, round_=True); add_rect(s, 7, ly + 6.5, 1.0, 8.5, ORANGE)
add_text(s, 9.5, ly + 6.5, 142, 8.5,
         [("", {"runs": [("Ablesbar:  ", {"bold": True, "color": ORANGE, "size": 11}),
             ("Viel Grün = harmonisierbar (Standard einführen). Rote/gelbe Häufungen (z. B. Fachbereich D) = "
              "Komplexitätstreiber → gezielt entscheiden: harmonisieren oder bewusst als Varianz zulassen und "
              "die Kosten (FTE, Konfig) transparent tragen.", {"color": DARK, "size": 11})]})], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)

# ==========================================================================
# SLIDE 5 -- Szenarien + Der Ask an den Vorstand
# ==========================================================================
s = prs.slides.add_slide(BLANK)
header(s, "ENTSCHEIDUNG", "Drei Szenarien — und was wir vom Vorstand brauchen", 5)
add_text(s, 7, 19.5, 146, 3.5, [("Der Trade-off, den der Vorstand entscheidet: Harmonisierung senkt dauerhafte FTE & Komplexität, kostet aber einmaligen Change.", {"size": 11, "color": GREY_TXT})], line_spacing=1.1)

# Szenarien-Tabelle
cols = [("Szenario", 34), ("Governance-Modell", 46), ("FTE (laufend)", 22), ("Komplexität", 22), ("Change (einmalig)", 22)]
tx, ty = 7, 24; hh = 5.5
add_rect(s, tx, ty, sum(w for _, w in cols), hh, DARK)
cxp = tx
for name, w in cols:
    add_text(s, cxp + 0.6, ty, w - 1, hh, [(name, {"size": 9.5, "bold": True, "color": WHITE})], anchor=MSO_ANCHOR.MIDDLE)
    cxp += w
rows = [
    ("A · Vollharmonisiert", "1 Standard, comply-or-explain", ("niedrig", GREEN), ("niedrig", GREEN), ("hoch", RED)),
    ("B · Föderiert (Empfehlung)", "Standard + wenige def. Varianten", ("mittel", AMBER), ("mittel", AMBER), ("mittel", AMBER)),
    ("C · Status quo digitalisiert", "je Fachbereich eigenes Modell", ("hoch", RED), ("hoch", RED), ("niedrig", GREEN)),
]
rh = 7.5
for i, (a, b, fte, kx, ch) in enumerate(rows):
    ry = ty + hh + i * rh
    add_rect(s, tx, ry, sum(w for _, w in cols), rh, WHITE if i % 2 == 0 else ROW_ALT)
    rec = (i == 1)
    if rec: add_rect(s, tx, ry, sum(w for _, w in cols), rh, GREEN_LT)
    add_rect(s, tx, ry, 0.7, rh, GREEN if rec else LINE_GREY)
    add_text(s, tx + 1.2, ry, 33, rh, [(a, {"size": 9.8, "bold": True, "color": DARK})], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, tx + 34.6, ry, 45, rh, [(b, {"size": 9, "color": GREY_TXT})], anchor=MSO_ANCHOR.MIDDLE)
    for k, (val, col) in enumerate([fte, kx, ch]):
        cxx = tx + 34 + 46 + k * 22
        add_text(s, cxx, ry, 22, rh, [(val, {"size": 9.5, "bold": True, "color": col, "align": PP_ALIGN.CENTER})], anchor=MSO_ANCHOR.MIDDLE)
add_rect(s, tx, ty, sum(w for _, w in cols), hh + len(rows) * rh, None, line=LINE_GREY, line_w=1.0)
add_text(s, 7, ty + hh + len(rows) * rh + 0.8, 146, 3,
         [("FTE-Zahlen werden im Folgeprojekt bottom-up je Governance-Rolle × Fachbereich × Standardisierungsgrad ermittelt.", {"size": 8.5, "italic": True, "color": LINE_GREY})])

# Der Ask
add_rect(s, 7, 60.5, 146, 22.5, GREEN_LT, round_=True); add_rect(s, 7, 60.5, 1.2, 22.5, GREEN)
add_text(s, 9.5, 61.3, 142, 4, [("Der Ask an den Vorstand", {"size": 13, "bold": True, "color": GREEN})])
asks = [
    ("Mandat & Budget", "Beauftragung des Folgeprojekts „Governance-Assessment & Zielprozess-Design“ (Bewertung Prozess + FTE)."),
    ("Grundsatzentscheid", "„Standard first / comply-or-explain“ als verbindliches Leitprinzip der Plattformeinführung."),
    ("Governance-Sponsor", "Benennung eines Vorstands-Sponsors, der Harmonisierung gegen Partikularinteressen absichert."),
    ("Entscheidungspunkt", "Szenario-Wahl (A/B/C) nach Vorliegen der Heatmap & FTE-Bewertung — als Gate vor dem Rollout."),
]
for i, (t, d) in enumerate(asks):
    lx = 9.5 + (i % 2) * 72; ly = 65.5 + (i // 2) * 8.5
    add_rect(s, lx, ly, 69, 7.5, WHITE, line=GREEN, line_w=0.75, round_=True)
    add_text(s, lx + 1.4, ly + 0.5, 66, 3, [(t, {"size": 10, "bold": True, "color": DARK})])
    add_text(s, lx + 1.4, ly + 3.6, 66, 4, [(d, {"size": 8.3, "color": GREY_TXT})], line_spacing=1.02)

prs.save(DST)
print("saved", DST, "with", len(prs.slides._sldIdLst), "slides")
