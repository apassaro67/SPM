"""Insert a new 'Business Outcomes' slide after current slide 11
('Vorprojekt SNOW – Ergebnisse 2 von 2'). The 9 ServiceNow business
outcomes from the uploaded deck are consolidated into 6 board-ready
rows in a clean table.
"""
from copy import deepcopy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v11.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v12.pptx"

ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
ORANGE_LT= RGBColor(0xFD, 0xE6, 0xCC)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
ROW_ALT  = RGBColor(0xEE, 0xF1, 0xF4)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)

U = 76200
def ex(u): return int(u * U)

prs = Presentation(SRC)

# Use the same layout as our current slide 9
layout = prs.slides[8].slide_layout
new_slide = prs.slides.add_slide(layout)
slide = new_slide

# Wipe placeholders
for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def add_oval(left, top, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK, default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(40000); tf.margin_right = Emu(40000)
    tf.margin_top = Emu(15000); tf.margin_bottom = Emu(15000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)
        if "space_after" in st:
            p.space_after = Pt(st["space_after"])

def add_text(left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.MIDDLE, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ========== HEADER ==========
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("Business Outcomes – Was SNOW SPM für STIHL konkret bedeutet",
     dict(size=20, bold=True, color=WHITE)),
], anchor=MSO_ANCHOR.TOP)
add_text(3, 3.2, 140, 2.4, [
    ("Die 9 ServiceNow Business Outcomes konsolidiert auf 6 wirkungsstarke STIHL-Ergebnisse",
     dict(size=12, italic=True, color=LIGHT_GREY)),
], anchor=MSO_ANCHOR.TOP)
add_rect(150, 0, 10, 6, ORANGE)
add_text(150, 0, 10, 6, [
    ("Outcomes", dict(size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
])

# ========== TABLE ==========
# Column geometry
COL_X = 4
TOTAL_W = 152
# # | Outcome | Ziel | STIHL-Wirkung
COL_W = [6, 30, 36, 80]   # sums to 152
def colx(i): return COL_X + sum(COL_W[:i])

# --- Table header row ---
hdr_y = 7.5
hdr_h = 4.5
add_rect(COL_X, hdr_y, TOTAL_W, hdr_h, DARK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
headers = [
    ("#",                     PP_ALIGN.CENTER),
    ("Business Outcome",      PP_ALIGN.LEFT),
    ("Was es bedeutet (Ziel)", PP_ALIGN.LEFT),
    ("STIHL-Wirkung – konkreter Impact",  PP_ALIGN.LEFT),
]
for i, (text, align) in enumerate(headers):
    pad_l = 1.5 if align == PP_ALIGN.LEFT else 0
    add_text(colx(i) + pad_l, hdr_y, COL_W[i] - pad_l, hdr_h, [
        (text, dict(size=11, bold=True, color=WHITE, align=align)),
    ], anchor=MSO_ANCHOR.MIDDLE)

# --- Data rows: 6 consolidated outcomes ---
rows = [
    ("Speed &\nTime-to-Market",
     "Schneller von Idee zur Markt-/Projektreife",
     "End-to-end-Digitalisierung Idee → MPM → VEW. Gates (SBGo, PH) im System. Workflows ersetzen wochenlange manuelle Schleifen zwischen Excel, SAP und PIT.",
     "konsolidiert: 1 Faster Time to Market + 7 Improved Speed Overall"),
    ("Resilience &\nAgility",
     "Schnell auf Störungen und Veränderungen reagieren",
     "Alternativ-Roadmaps ohne neue Excel-Iterationen. Bei Budgetkürzung oder Marktveränderung sofort Re-Priorisierung – Effekte auf Kapazität und Strategie sind unmittelbar sichtbar.",
     "konsolidiert: 2 Response to Disruptions + 6 Increased Agility"),
    ("Strategic Alignment\n& Transparency",
     "Strategie und Umsetzung durchgängig verknüpfen",
     "OKRs/Ziele direkt mit Programmen, Projekten und Demands verbunden. Vorstands-Cockpit mit Ampelstatus, Budget-Auslastung und OKR-Beitrag aller Projekte.",
     "konsolidiert: 3 Strategy ↔ Execution + 8 Big-picture Focus"),
    ("Efficiency",
     "Doppelarbeit und manuelle Reportingschleifen vermeiden",
     "Priorisierung über Scoring-Logik (Market / Business Case) ersetzt Excel-Bewertungen. Stunden im System, SAP-Ist-Kosten automatisch eingebunden – Plan-vs-Ist live, ohne manuelle Datenladevorgänge.",
     "Quelle: 4 Improved Efficiency"),
    ("Cross-Functional\nCohesion",
     "Eine Plattform für alle Disziplinen",
     "Einheitliches Datenmodell und Templates für MPM, VEW/PEP, VPM und IT. Cross-funktionale Programme mit gemeinsamer Freigabe, Abhängigkeiten und Reporting – Insellösungen entfallen.",
     "Quelle: 5 Cohesion of Multiple Disciplines"),
    ("Business Value\nRealization",
     "Nutzen messbar machen und nachverfolgen",
     "Business Cases strukturiert erfasst und über den gesamten Projektlebenszyklus nachverfolgt. Nach Projektabschluss Soll-/Ist-Abgleich gegen Strategie- und Nutzen-Ziele.",
     "Quelle: 9 Business Value Realization"),
]

row_top = hdr_y + hdr_h + 0.5
row_h   = 11.0
gap     = 0.6

for i, (title, ziel, impact, src) in enumerate(rows):
    y = row_top + i * (row_h + gap)
    bg = LIGHT_BG if i % 2 == 0 else ROW_ALT
    add_rect(COL_X, y, TOTAL_W, row_h, bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # Left orange accent
    add_rect(COL_X, y, 0.7, row_h, ORANGE)

    # Number badge
    add_oval(colx(0) + 1.3, y + row_h / 2 - 1.6, 3.2, 3.2, ORANGE)
    add_text(colx(0) + 1.3, y + row_h / 2 - 1.6, 3.2, 3.2, [
        (str(i + 1), dict(size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)

    # Outcome title (column 1)
    add_text(colx(1) + 0.4, y + 0.4, COL_W[1] - 0.8, row_h - 3.5, [
        (title.replace("\n", " "),
         dict(size=12.5, bold=True, color=DARK)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    # Source / consolidation note in italic small
    add_text(colx(1) + 0.4, y + row_h - 3.0, COL_W[1] - 0.8, 2.6, [
        (src, dict(size=8, italic=True, color=GREY_TXT)),
    ], anchor=MSO_ANCHOR.TOP)

    # Ziel (column 2)
    add_text(colx(2) + 0.4, y + 0.4, COL_W[2] - 0.8, row_h - 0.8, [
        (ziel, dict(size=11, color=DARK)),
    ], anchor=MSO_ANCHOR.MIDDLE)

    # STIHL-Wirkung (column 3)
    add_text(colx(3) + 0.6, y + 0.4, COL_W[3] - 1.2, row_h - 0.8, [
        (impact, dict(size=10.5, color=DARK)),
    ], anchor=MSO_ANCHOR.MIDDLE)

# ========== FOOTER ==========
fy = row_top + 6 * (row_h + gap) + 0.4
add_text(COL_X, fy, TOTAL_W, 2.0, [
    ("Konsolidiert aus 'Stihl_SPM_Business_Outcomes' – 9 ursprüngliche Outcomes (Faster Time to Market, Response to Disruptions, Strategy/Execution, Efficiency, Cohesion, Agility, Speed Overall, Big-picture, Value Realization).",
     dict(size=8.5, italic=True, color=GREY_TXT)),
])
add_rect(0, 89.3, 160, 0.7, ORANGE)

# ========== Move new slide right after current slide 11 (index 10) ==========
sldIdLst = prs.slides._sldIdLst
slides = list(sldIdLst)
new_id = slides[-1]
sldIdLst.remove(new_id)
# Insert at position 11 (after current slide 11 which is at index 10)
sldIdLst.insert(11, new_id)

prs.save(DST)
print("Saved:", DST)
