"""Redesigned slide 9: compact ServiceNow capabilities (top), reference
customers with branded logo tiles (middle), STIHL-specific info (bottom).
No more 'Kernbotschaften' bullets at the bottom.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v5.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v6.pptx"

ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)

U = 76200
def ex(u): return int(u * U)

prs = Presentation(SRC)
slide = prs.slides[8]  # the slide we just created

# Wipe everything
for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.2):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK, default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(50000); tf.margin_right = Emu(50000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)
        if "space_after" in st:
            p.space_after = Pt(st["space_after"])

def add_text(left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.TOP, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ============== HEADER ==============
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("ServiceNow als Tool – Reichweite, Referenzkunden und STIHL-Kontext",
     dict(size=20, bold=True, color=WHITE)),
])
add_text(3, 3.2, 140, 2.4, [
    ("Bewährter Industriestandard, große deutsche Referenzbasis und seit 2014 bei STIHL etabliert",
     dict(size=12, italic=True, color=LIGHT_GREY)),
])
add_rect(150, 0, 10, 6, ORANGE)
tb = slide.shapes.add_textbox(ex(150), ex(0), ex(10), ex(6))
tb.fill.background(); tb.line.fill.background()
tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
set_paragraphs(tb, [("Folie 9", dict(size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER))])

# ============== TOP THIRD: ServiceNow Capabilities (compact) ==============
top_y = 8
add_text(4, top_y, 152, 2.4, [
    ("Was ServiceNow als Tool bietet",
     dict(size=13, bold=True, color=DARK)),
])

# Subline
add_text(4, top_y + 2.2, 152, 1.8, [
    ("Cloud-Plattform mit einer Datenbasis – Workflows, KI (Now Assist), Low-Code & Integration für das gesamte Unternehmen",
     dict(size=10.5, italic=True, color=GREY_TXT)),
])

# 5 capability pills (ITSM/ITOM, SPM, HR, CSM, Security/GRC)
pills = [
    ("IT Service Management\n& IT Operations", "ITSM · ITOM"),
    ("Strategic Portfolio\nManagement (SPM)", "Strategy → Execute"),
    ("HR Service Delivery", "Employee Workflows"),
    ("Customer Service\nManagement", "Customer Workflows"),
    ("Security Operations\n& GRC", "Risk · Compliance"),
]
pill_top = 12.2
pill_h = 7.0
gap = 1.0
total = 152
pw = (total - 4 * gap) / 5
for i, (head, sub) in enumerate(pills):
    x = 4 + i * (pw + gap)
    is_spm = "SPM" in head
    fill = ORANGE if is_spm else WHITE
    text_color = WHITE if is_spm else DARK
    sub_color = LIGHT_GREY if is_spm else GREY_TXT
    add_rect(x, pill_top, pw, pill_h, fill, line=ORANGE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line_w=1.5 if is_spm else 1.0)
    add_text(x, pill_top + 0.6, pw, pill_h - 1.6, [
        (head.split("\n")[0], dict(size=10.5, bold=True, color=text_color, align=PP_ALIGN.CENTER)),
        (head.split("\n")[1] if "\n" in head else "",
         dict(size=10.5, bold=True, color=text_color, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    add_text(x, pill_top + pill_h - 2.0, pw, 1.6, [
        (sub, dict(size=8.8, italic=True, color=sub_color, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)

# Foundation row (key building blocks)
fy = pill_top + pill_h + 1.0
add_rect(4, fy, total, 4.0, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
add_text(5, fy, total - 2, 4.0, [
    ("Plattform-Bausteine:  Workflow Engine  ·  Now Assist (KI)  ·  Low-Code Studio  ·  Integration Hub  ·  Performance Analytics  ·  CMDB",
     dict(size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
], anchor=MSO_ANCHOR.MIDDLE)

# ============== MIDDLE THIRD: Reference Customers ==============
ref_top = 30.5
add_text(4, ref_top, total, 2.4, [
    ("Große deutsche Referenzkunden – ServiceNow im Einsatz in IT und Geschäftsprozessen",
     dict(size=13, bold=True, color=DARK)),
])

# Brand-colored typographic logo tiles
customers = [
    ("SIEMENS",          "Industrie / Tech",       RGBColor(0x00, 0x9D, 0x9D), WHITE),
    ("SAP",              "Software",               RGBColor(0x00, 0x3A, 0x5D), RGBColor(0x07, 0xC8, 0xF0)),
    ("Allianz",          "Versicherung",           RGBColor(0x00, 0x37, 0x81), WHITE),
    ("Deutsche Bank",    "Banking",                RGBColor(0x00, 0x18, 0xA8), WHITE),
    ("T",                "Deutsche Telekom",       RGBColor(0xE2, 0x00, 0x74), WHITE),
    ("BOSCH",            "Industrie / Automotive", RGBColor(0xC8, 0x10, 0x2E), WHITE),
    ("Mercedes-Benz",    "Automotive",             RGBColor(0x00, 0x00, 0x00), WHITE),
    ("Volkswagen",       "Automotive",             RGBColor(0x00, 0x1E, 0x50), WHITE),
]

box_top = 33.5
box_h = 11.5
n = len(customers)
gap_c = 0.8
box_w = (total - (n - 1) * gap_c) / n

for i, (name, branche, bg, fg) in enumerate(customers):
    x = 4 + i * (box_w + gap_c)
    # Logo tile
    add_rect(x, box_top, box_w, box_h - 4, bg,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # Logo text (large, brand-style)
    name_size = 16 if len(name) <= 4 else (13 if len(name) <= 8 else 11)
    add_text(x, box_top, box_w, box_h - 4, [
        (name, dict(size=name_size, bold=True, color=fg, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    # Caption: Branche
    add_text(x, box_top + box_h - 3.6, box_w, 3.4, [
        (branche, dict(size=8.5, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.TOP)

# Footnote logo
add_text(4, 46.5, total, 2.0, [
    ("Stilisierte Markenfarben als Logo-Platzhalter – echte Logos können in der finalen Vorstandsfassung eingesetzt werden  |  Quelle: ServiceNow Customer References, öffentliche Quellen",
     dict(size=8, italic=True, color=LIGHT_GREY, align=PP_ALIGN.LEFT)),
])

# Divider
add_rect(4, 49.5, total, 0.15, ORANGE)

# ============== BOTTOM THIRD: STIHL-spezifisch ==============
bot_top = 51.0
add_text(4, bot_top, total, 2.4, [
    ("ServiceNow @ STIHL – die Plattform ist bereits etabliert",
     dict(size=14, bold=True, color=DARK)),
])

stihl_boxes = [
    ("Im Einsatz seit 2014",
     "Globaler IT Service Management-Standard – inklusive China.\n"
     "Etabliertes Betriebsteam, eingespielte Lieferantenbeziehung.",
     "12 Jahre Betrieb"),
    ("DSGVO-konformes Hosting in Deutschland",
     "Hosting in den ServiceNow-Datacentern Frankfurt am Main und Düsseldorf.\n"
     "Vertragliche und technische DSGVO-Konformität gegeben.",
     "Frankfurt + Düsseldorf"),
    ("Anschluss-/Erweiterungsfähigkeit",
     "ITSM-CMDB als Andockpunkt für SPM (Apps, Services, Verträge).\n"
     "Plattform-Erweiterung statt Tool-Neubeschaffung – Synergien bei Lizenz, Betrieb, Schulung.",
     "Plattform-Synergie"),
]

sb_top = 53.5
sb_h = 21
gap_s = 2
sw = (total - 2 * gap_s) / 3
for i, (head, body, badge) in enumerate(stihl_boxes):
    x = 4 + i * (sw + gap_s)
    add_rect(x, sb_top, sw, sb_h, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # accent strip
    add_rect(x, sb_top, sw, 1.0, GREEN, shape=MSO_SHAPE.RECTANGLE)
    # Badge top right
    badge_w = 16
    add_rect(x + sw - badge_w - 0.8, sb_top + 1.6, badge_w, 2.4, GREEN,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x + sw - badge_w - 0.8, sb_top + 1.6, badge_w, 2.4, [
        (badge, dict(size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    # Headline
    add_text(x + 1.5, sb_top + 1.8, sw - badge_w - 3.0, 4.5, [
        (head, dict(size=12, bold=True, color=DARK)),
    ])
    # Body
    body_paras = [(line, dict(size=10, color=DARK, space_after=4))
                  for line in body.split("\n") if line]
    add_text(x + 1.5, sb_top + 7.0, sw - 3, sb_h - 8, body_paras)

# ============== FOOTER ==============
add_text(4, 86.3, total, 2.2, [
    ("Quellen: ServiceNow Produkt- und Kunden-Referenzen (servicenow.com/customers); STIHL IT (Bestand seit 2014, Hosting Frankfurt + Düsseldorf, DSGVO-konform).",
     dict(size=8.5, italic=True, color=GREY_TXT)),
])
add_rect(0, 89.3, 160, 0.7, ORANGE)

prs.save(DST)
print("Saved:", DST)
