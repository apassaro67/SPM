"""Slide 10 v22: ServiceNow POC @ STIHL - placeholder layout for screenshots.

Four picture placeholders (Prioritization with Lens, Product Roadmap,
Goals & Targets, Strategic Hierarchy detail) with descriptive captions.
User drops actual screenshots into the placeholders in PowerPoint.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v21.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v22.pptx"

ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)
DASH_GREY = RGBColor(0xA8, 0xAE, 0xB6)
NOW_GREEN_DK = RGBColor(0x29, 0x9B, 0x32)

U = 76200
def ex(u): return int(u * U)

prs = Presentation(SRC)
slide = prs.slides[9]

# Wipe
for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE,
             line_w=1.0, dashed=False):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
        if dashed:
            ln = s.line._get_or_add_ln()
            prstDash = etree.SubElement(ln, qn('a:prstDash'))
            prstDash.set('val', 'dash')
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK,
                   default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(20000); tf.margin_right = Emu(20000)
    tf.margin_top = Emu(10000); tf.margin_bottom = Emu(10000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)

def add_text(left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.MIDDLE, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ============== HEADER ==============
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("ServiceNow SPM PoC bei STIHL – Beispiele aus dem Strategic Planning Workspace",
     dict(size=18, bold=True, color=WHITE))], anchor=MSO_ANCHOR.TOP)
add_text(3, 3.2, 140, 2.4, [
    ("Strategien, Portfolios, Roadmaps und Ziele live auf der Plattform – STIHL TRAIL 2030 Demo",
     dict(size=12, italic=True, color=LIGHT_GREY))], anchor=MSO_ANCHOR.TOP)
add_rect(150, 0, 10, 6, ORANGE)
add_text(150, 0, 10, 6, [
    ("Folie 10", dict(size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER))])

# ============== 2x2 GRID OF PICTURE PLACEHOLDERS ==============
GRID_X = 4
GRID_Y = 7.5
GRID_W = 152
GRID_H = 76

GAP = 1.5
CARD_W = (GRID_W - GAP) / 2
CARD_H = (GRID_H - GAP) / 2 - 2.0   # leave room for slide footer

cards = [
    ("Prioritization mit Strategic-Investments-Lens",
     "STIHL TRAIL 2030 Demo – Lens-Hierarchie zeigt Strategic Priorities, Initiatives, Programs, Projects und Demands mit Gantt-Visualisierung. Filter über Lens-Auswahl (Org Unit, Strategic Investments, …).",
     "Bild 1 / 4"),
    ("Goals & Targets mit Status und Progress",
     "OKR-/Ziel-Hierarchie mit Status-Ampel (Grün/Gelb/Rot), Fortschritt in % und Trend-Charts pro Ziel. Side-Panel mit Soll-/Ist-Vergleich.",
     "Bild 2 / 4"),
    ("Stihl Product Roadmap – Swim-Lane-Planung",
     "Portfolio-Planning-Workbench mit Produkt-Swim-Lanes (z. B. PMS 180.0, WKP 250.0, ZMS 400.0). FY26-Quartale + Monatsraster; Demands und Projekte direkt auf der Zeitachse.",
     "Bild 3 / 4"),
    ("Strategic Hierarchy – Programme & Initiativen",
     "Volle Lens-Hierarchie: STIHL TRAIL 2030 → Customer-Centric Products → Connected → einzelne Vorhaben (Stihl Global AI Hub, EURO Umstellung Bulgarien, OT Security, EMEA Shared Datacenter …).",
     "Bild 4 / 4"),
]

for i, (head, caption, badge) in enumerate(cards):
    col = i % 2
    row = i // 2
    x = GRID_X + col * (CARD_W + GAP)
    y = GRID_Y + row * (CARD_H + GAP)

    # Card frame (light background)
    add_rect(x, y, CARD_W, CARD_H, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

    # Header band
    add_rect(x, y, CARD_W, 4.0, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x + 1.5, y + 0.3, CARD_W - 14, 3.4, [
        (head, dict(size=11.5, bold=True, color=WHITE))],
        anchor=MSO_ANCHOR.MIDDLE)
    # Badge
    add_rect(x + CARD_W - 11.5, y + 0.7, 10, 2.6, ORANGE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x + CARD_W - 11.5, y + 0.7, 10, 2.6, [
        (badge, dict(size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER))],
        anchor=MSO_ANCHOR.MIDDLE)

    # Picture placeholder area (dashed border, white background)
    pic_x = x + 1.5
    pic_y = y + 4.8
    pic_w = CARD_W - 3.0
    pic_h = CARD_H - 4.8 - 6.0   # leave space for caption
    add_rect(pic_x, pic_y, pic_w, pic_h, WHITE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line=DASH_GREY, line_w=1.5, dashed=True)
    # Camera icon hint + instruction
    add_text(pic_x, pic_y, pic_w, pic_h, [
        ("📷  Hier Screenshot einfügen",
         dict(size=14, bold=True, color=DASH_GREY, align=PP_ALIGN.CENTER)),
        ("Drag-and-Drop in PowerPoint oder Einfügen → Bild …",
         dict(size=9, italic=True, color=DASH_GREY, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)

    # Caption below image
    cap_y = pic_y + pic_h + 0.4
    cap_h = CARD_H - (cap_y - y) - 0.6
    add_text(x + 1.5, cap_y, CARD_W - 3.0, cap_h, [
        (caption, dict(size=9.5, color=DARK))],
        anchor=MSO_ANCHOR.TOP)

# ============== FOOTER ==============
add_rect(0, 89.3, 160, 0.7, ORANGE)
add_text(4, 86.3, 152, 2.0, [
    ("PoC-Screenshots aus ServiceNow Strategic Planning Workspace / Portfolio Planning – Demo-Instance STIHL TRAIL 2030.",
     dict(size=8.5, italic=True, color=GREY_TXT))])

prs.save(DST)
print("Saved:", DST)
