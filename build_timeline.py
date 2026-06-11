"""New slide v23: Detailed timeline 18.05.2026 - 31.10.2026
showing the activities between board meeting and Q3 board meeting.
Inserted right after current slide 12 ('Weiteres Vorgehen').
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from datetime import date

SRC = "/home/user/SPM/2026_05_18_VS_Zielbild_SNOW_SPM_v22.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v23.pptx"

# Palette
ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
ORANGE_LT= RGBColor(0xFD, 0xE6, 0xCC)
ORANGE2  = RGBColor(0xE9, 0x6C, 0x0C)
ORANGE3  = RGBColor(0xC9, 0x53, 0x00)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
BLUE_HD  = RGBColor(0x3B, 0x6B, 0xA5)
PURPLE   = RGBColor(0x86, 0x5D, 0xE4)
RED_HL   = RGBColor(0xB9, 0x1C, 0x1C)
GREY_LN  = RGBColor(0xCB, 0xD2, 0xD9)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)

U = 76200
def ex(u): return int(u * U)

prs = Presentation(SRC)
# Use layout of an existing slide for consistency
layout = prs.slides[9].slide_layout
new_slide = prs.slides.add_slide(layout)
slide = new_slide
for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def add_diamond(left, top, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK,
                   default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(20000); tf.margin_right = Emu(20000)
    tf.margin_top = Emu(10000); tf.margin_bottom = Emu(10000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)

def add_text(left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.MIDDLE, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ========== HEADER ==========
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("Detail-Terminierung 18.05. – 31.10.2026",
     dict(size=20, bold=True, color=WHITE))], anchor=MSO_ANCHOR.TOP)
add_text(3, 3.2, 140, 2.4, [
    ("Vorbereitende Aktivitäten zwischen Vorstandstermin und Vorstandssitzung Q3 2026",
     dict(size=12, italic=True, color=LIGHT_GREY))], anchor=MSO_ANCHOR.TOP)
add_rect(150, 0, 10, 6, ORANGE)
add_text(150, 0, 10, 6, [
    ("Terminplan", dict(size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER))])

# ========== TIMELINE GEOMETRY ==========
# Time range: 18.05.2026 to 31.10.2026 (167 days)
START = date(2026, 5, 18)
END   = date(2026, 10, 31)
TOTAL_DAYS = (END - START).days  # 166
def day_index(d): return (d - START).days
def frac(d): return day_index(d) / TOTAL_DAYS

LANE_LBL_X = 4
LANE_LBL_W = 36
TL_X = LANE_LBL_X + LANE_LBL_W + 1.0
TL_RIGHT = 156
TL_W = TL_RIGHT - TL_X
def dx(d): return TL_X + frac(d) * TL_W

# ========== MONTH AXIS ==========
axis_top = 7.5
month_h = 3.6
months = [
    (date(2026, 5, 18), date(2026, 5, 31), "Mai 26",  ORANGE_LT),
    (date(2026, 6, 1),  date(2026, 6, 30), "Juni 26", LIGHT_BG),
    (date(2026, 7, 1),  date(2026, 7, 31), "Juli 26", ORANGE_LT),
    (date(2026, 8, 1),  date(2026, 8, 31), "Aug 26",  LIGHT_BG),
    (date(2026, 9, 1),  date(2026, 9, 30), "Sept 26", ORANGE_LT),
    (date(2026, 10,1),  date(2026, 10,31), "Okt 26",  LIGHT_BG),
]
add_text(LANE_LBL_X, axis_top, LANE_LBL_W, month_h, [
    ("ZEITRAUM ▶", dict(size=11, bold=True, color=DARK, align=PP_ALIGN.RIGHT))],
    anchor=MSO_ANCHOR.MIDDLE)
for s_d, e_d, label, col in months:
    x0 = dx(s_d)
    x1 = dx(e_d)
    add_rect(x0, axis_top, x1 - x0 - 0.05, month_h, col,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x0, axis_top, x1 - x0 - 0.05, month_h, [
        (label, dict(size=11, bold=True, color=DARK, align=PP_ALIGN.CENTER))],
        anchor=MSO_ANCHOR.MIDDLE)

# Month-boundary tick marks below the axis
tick_top = axis_top + month_h + 0.1
for d in [date(2026,6,1), date(2026,7,1), date(2026,8,1), date(2026,9,1), date(2026,10,1)]:
    x = dx(d)
    add_rect(x - 0.04, tick_top, 0.08, 0.6, GREY_LN)

# ========== MILESTONE ROW ==========
ms_top = axis_top + month_h + 1.4
ms_h = 9.0
add_text(LANE_LBL_X, ms_top, LANE_LBL_W, ms_h, [
    ("MEILENSTEINE", dict(size=11, bold=True, color=DARK, align=PP_ALIGN.RIGHT))],
    anchor=MSO_ANCHOR.MIDDLE)
# horizontal line
line_y = ms_top + ms_h / 2
add_rect(TL_X, line_y - 0.05, TL_W, 0.1, GREY_LN)

milestones = [
    (date(2026, 5, 18), "VS 18.05.2026", "Vorstand – heutige Sitzung", RED_HL, True,  "above"),
    (date(2026, 9, 25), "VS Ende Sept.", "Vorstand Q3 2026",           RED_HL, True,  "below"),
]
LBL_W = 22
for d, head, sub, col, big, side in milestones:
    cx = dx(d)
    diam = 2.4 if big else 1.6
    add_diamond(cx - diam/2, line_y - diam/2, diam, diam, col)
    if side == "above":
        add_rect(cx - 0.04, ms_top + 0.5, 0.08,
                 (line_y - diam/2) - (ms_top + 0.5), GREY_LN)
        add_text(cx - LBL_W/2, ms_top - 0.2, LBL_W, 2.0, [
            (head, dict(size=10.5, bold=True, color=col, align=PP_ALIGN.CENTER))],
            anchor=MSO_ANCHOR.BOTTOM)
        add_text(cx - LBL_W/2, ms_top + 1.8, LBL_W, 1.8, [
            (sub, dict(size=8.5, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER))],
            anchor=MSO_ANCHOR.MIDDLE)
    else:
        add_rect(cx - 0.04, line_y + diam/2,
                 0.08, (ms_top + ms_h - 0.5) - (line_y + diam/2), GREY_LN)
        add_text(cx - LBL_W/2, ms_top + ms_h - 2.4, LBL_W, 1.8, [
            (head, dict(size=10.5, bold=True, color=col, align=PP_ALIGN.CENTER))],
            anchor=MSO_ANCHOR.BOTTOM)
        add_text(cx - LBL_W/2, ms_top + ms_h - 0.6, LBL_W, 1.6, [
            (sub, dict(size=8.5, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER))],
            anchor=MSO_ANCHOR.MIDDLE)

# ========== SWIM LANES ==========
# (Label, sublabel, start, end, color, bar_label)
lanes = [
    ("Prozess- & Toolanalyse",
     "Planhorizon · Anforderungen",
     date(2026, 5, 30), date(2026, 7, 15),
     ORANGE,
     "Analyse · Klärung · Planhorizon"),
    ("Business Case",
     "Wirtschaftlichkeit, ROI",
     date(2026, 6, 15), date(2026, 8, 30),
     BLUE_HD,
     "Business Case Ausarbeitung"),
    ("Alternative PM-Plattform",
     "VEW / VPM (Hedge)",
     date(2026, 6, 11), date(2026, 8, 30),
     NAVY,
     "Bewertung alternativer PM-Tools"),
    ("Stakeholder-Management",
     "Vorstandsinfos je Ressort",
     date(2026, 6, 30), date(2026, 10, 16),
     ORANGE2,
     "Ressort-Bilateralen · Abstimmungs-Workshops"),
    ("Budgetplanung",
     "Plan- & Investitions-Mittel",
     date(2026, 8, 1), date(2026, 8, 30),
     GREEN,
     "Budgetplanung"),
    ("RFP Dienstleisterauswahl",
     "Ausschreibung & Vergabe",
     date(2026, 10, 1), date(2026, 10, 31),
     ORANGE3,
     "RFP · Bewertung · Vergabe"),
]
lane_top0 = ms_top + ms_h + 1.0
lane_h = 7.5
lane_gap = 0.6

for i, (label, sublabel, sd, ed, color, bar_lbl) in enumerate(lanes):
    y = lane_top0 + i * (lane_h + lane_gap)
    # Label cell
    add_rect(LANE_LBL_X, y, LANE_LBL_W, lane_h, LIGHT_BG,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(LANE_LBL_X, y, 0.7, lane_h, color)
    add_text(LANE_LBL_X + 1.4, y + 0.4, LANE_LBL_W - 2, 3.0, [
        (label, dict(size=11, bold=True, color=DARK))],
        anchor=MSO_ANCHOR.TOP)
    add_text(LANE_LBL_X + 1.4, y + 3.4, LANE_LBL_W - 14, 3.2, [
        (sublabel, dict(size=8.8, italic=True, color=GREY_TXT))],
        anchor=MSO_ANCHOR.TOP)
    # Date range badge top right
    add_rect(LANE_LBL_X + LANE_LBL_W - 13, y + 0.5, 12, 2.6, color,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(LANE_LBL_X + LANE_LBL_W - 13, y + 0.5, 12, 2.6, [
        (f"{sd.strftime('%d.%m.')} – {ed.strftime('%d.%m.')}",
         dict(size=8.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER))],
        anchor=MSO_ANCHOR.MIDDLE)

    # Lane background
    add_rect(TL_X, y, TL_W, lane_h, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # Month boundary gridlines
    for d in [date(2026,6,1), date(2026,7,1), date(2026,8,1),
              date(2026,9,1), date(2026,10,1)]:
        gx = dx(d)
        add_rect(gx - 0.04, y + 0.4, 0.08, lane_h - 0.8, GREY_LN)

    # Activity bar
    bar_x = dx(sd)
    bar_w = dx(ed) - bar_x - 0.1
    bar_y = y + 1.0
    bar_h = lane_h - 2.0
    add_rect(bar_x, bar_y, bar_w, bar_h, color,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    title_size = 11 if bar_w >= 20 else (10 if bar_w >= 14 else 9)
    add_text(bar_x + 0.4, bar_y, bar_w - 0.8, bar_h, [
        (bar_lbl,
         dict(size=title_size, bold=True, color=WHITE, align=PP_ALIGN.CENTER))],
        anchor=MSO_ANCHOR.MIDDLE)

# Today marker (vertical line at 18.05)
total_lane_zone = len(lanes) * (lane_h + lane_gap)
today_x = dx(date(2026, 5, 18))
add_rect(today_x - 0.06, lane_top0 - 0.3, 0.12, total_lane_zone + 0.4, RED_HL)
# VS Ende Sept marker
vs_x = dx(date(2026, 9, 25))
add_rect(vs_x - 0.05, lane_top0 - 0.3, 0.10, total_lane_zone + 0.4, RED_HL)

# ========== LEGEND / FOOTER ==========
leg_top = lane_top0 + total_lane_zone + 0.5
add_text(LANE_LBL_X, leg_top, 30, 1.8, [
    ("Legende:", dict(size=9, bold=True, color=DARK))], anchor=MSO_ANCHOR.MIDDLE)
items = [
    (ORANGE,   "Analyse"),
    (BLUE_HD,  "Business Case"),
    (NAVY,     "Alternative"),
    (ORANGE2,  "Stakeholder"),
    (GREEN,    "Budget"),
    (ORANGE3,  "RFP"),
    (RED_HL,   "◆ Vorstands-Meilenstein"),
]
cx = LANE_LBL_X + 14
for col, label in items:
    if "◆" in label:
        add_diamond(cx, leg_top + 0.6, 1.4, 1.4, col)
        add_text(cx + 1.7, leg_top, 24, 1.8, [
            (label, dict(size=8.5, color=DARK))], anchor=MSO_ANCHOR.MIDDLE)
        cx += 28
    else:
        add_rect(cx, leg_top + 0.7, 1.4, 1.4, col,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        add_text(cx + 1.7, leg_top, 18, 1.8, [
            (label, dict(size=8.5, color=DARK))], anchor=MSO_ANCHOR.MIDDLE)
        cx += 19

# Footer
add_text(4, 86.3, 152, 2.0, [
    ("Detail-Terminierung zwischen Vorstand 18.05.2026 und Vorstand Q3 2026.  Stand: 18.05.2026.",
     dict(size=8.5, italic=True, color=GREY_TXT))])
add_rect(0, 89.3, 160, 0.7, ORANGE)

# ========== Insert as new slide after current slide 12 ==========
sldIdLst = prs.slides._sldIdLst
slides = list(sldIdLst)
new_id = slides[-1]
sldIdLst.remove(new_id)
# Insert at index 12 (right after current slide 12 'Weiteres Vorgehen' at index 11)
sldIdLst.insert(12, new_id)

prs.save(DST)
print("Saved:", DST)
