"""Build Word-Bericht: ServiceNow SPM Pro vs. Planview vs. Planisware."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

DST = "/tmp/2026_10_15_PPM_Vergleich_SNOW_Planview_Planisware.docx"

# STIHL colors
ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
RED_HL   = RGBColor(0xB9, 0x1C, 0x1C)
LIGHT_BG_HEX = "F7F7F7"
ORANGE_HEX   = "F07F12"
NAVY_HEX     = "2C3E50"
GREEN_HEX    = "2E7D32"
ORANGE_LT_HEX= "FDE6CC"
GREY_BG_HEX  = "E9EDF1"

doc = Document()

# ----- Page margins -----
for section in doc.sections:
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# ----- Default style -----
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = DARK

# ----- Helpers -----
def add_heading(text, level=1, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.bold = True
    r.font.size = Pt({1: 20, 2: 15, 3: 12}.get(level, 12))
    r.font.color.rgb = color or (ORANGE if level == 1 else NAVY)
    return p

def add_para(text, italic=False, bold=False, color=None, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.italic = italic
    r.bold = bold
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return p

def add_bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(it)
        r.font.name = "Calibri"
        r.font.size = Pt(11)

def cell_shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def cell_set_text(cell, text, *, bold=False, color=None, size=10.5, align=None,
                  hex_fill=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if hex_fill:
        cell_shade(cell, hex_fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def set_cell_border(cell, color_hex="D1D5DB", size="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), size)
        b.set(qn("w:color"), color_hex)
        tcBorders.append(b)
    tcPr.append(tcBorders)

def build_table(headers, rows, col_widths_cm=None, header_fill=NAVY_HEX,
                first_col_fill=None, first_col_bold=True, alt_row_fill=True,
                center_cols=None):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for c in range(len(rows) + 1):
                tbl.cell(c, i).width = Cm(w)
    # Header
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        cell_set_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                      size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, hex_fill=header_fill)
        set_cell_border(c)
    # Data
    for r, row in enumerate(rows, start=1):
        for i, v in enumerate(row):
            c = tbl.cell(r, i)
            fill = None
            if alt_row_fill and r % 2 == 0:
                fill = "EEF1F4"
            elif alt_row_fill:
                fill = LIGHT_BG_HEX
            if i == 0 and first_col_fill:
                fill = first_col_fill
            align = None
            if center_cols and i in center_cols:
                align = WD_ALIGN_PARAGRAPH.CENTER
            cell_set_text(c, v, bold=(i == 0 and first_col_bold),
                          size=10.5, align=align, hex_fill=fill)
            set_cell_border(c)
    return tbl

def add_callout(text, color_hex, title=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Cm(17)
    c = tbl.cell(0, 0)
    c.text = ""
    if title:
        p = c.paragraphs[0]
        r = p.add_run(title + "\n")
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(12)
    p2 = c.add_paragraph()
    r2 = p2.add_run(text)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    cell_shade(c, color_hex)
    set_cell_border(c, color_hex="808080", size="2")
    doc.add_paragraph()

# ============================================================
# TITLE + META
# ============================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("ServiceNow SPM Pro vs. Planview vs. Planisware")
r.bold = True
r.font.name = "Calibri"
r.font.size = Pt(24)
r.font.color.rgb = ORANGE

p = doc.add_paragraph()
r = p.add_run("Vergleich und Empfehlung im Kontext einer bestehenden ServiceNow-Basis")
r.font.name = "Calibri"
r.font.size = Pt(13)
r.font.color.rgb = GREY_TXT
r.italic = True

# Meta line
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
meta = p.add_run("Kontext: Gartner Magic Quadrant für Adaptive Project Management and Reporting (Juni 2026)  ·  "
                 "Fokus: Projekt-Management-Fähigkeiten  ·  Empfehlung für Migration von altem PM-Tool")
meta.font.size = Pt(10)
meta.font.color.rgb = GREY_TXT
meta.italic = True

# Horizontal rule
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "8")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), ORANGE_HEX)
pBdr.append(bottom)
pPr.append(pBdr)

# ============================================================
# EXECUTIVE SUMMARY
# ============================================================
add_heading("Executive Summary", level=1)

add_para(
    "Die Gartner-Leader Planview und Planisware sind im reinen Projekt-Management (Gantt, "
    "Ressourcen, Finanzen) auf höchstem Niveau, aber Insellösungen. ServiceNow SPM Pro liegt "
    "im Projekt-Management etwas dahinter, gewinnt jedoch durch Plattform-Integration, KI und "
    "Extensibility deutlich an Boden.",
    space_after=8
)
add_para(
    "Für ein Unternehmen mit bestehender ServiceNow-Basis und aktivem SPM Demand-Management "
    "ist die klare Empfehlung: ServiceNow SPM Pro als Ziel-Plattform, ergänzt durch App Engine "
    "für spezifische Prozesse (z.B. Stage-Gate-PEP). Ein hybrider Ansatz mit Planisware im "
    "R&D-Bereich bleibt als Hedge denkbar, ist aber wirtschaftlich meist unterlegen.",
    space_after=12
)

add_callout(
    "ServiceNow SPM Pro empfohlen — Datenkontinuität, Betriebskontinuität und Time-to-Value "
    "überwiegen deutlich die kleinen Funktionsvorteile der spezialisierten PM-Tools.",
    color_hex="DDF1DE",
    title="🎯 Kernempfehlung"
)

# ============================================================
# STECKBRIEFE
# ============================================================
add_heading("Steckbriefe der drei Player", level=1)

# Planview
add_heading('🏆 Planview — Leader mit höchster „Ability to Execute"', level=2)
add_bullets([
    "Historisch stärkster Enterprise-PPM-Player (Ex-Planview Enterprise One, AdaptiveWork)",
    "Portfolio: Planview Portfolios, AdaptiveWork (Ex-Clarizen), LeanKit (Kanban), ProjectPlace",
    "Kernstärken: Ressourcen-Optimierung, Kapazitätsplanung, Multi-Portfolio-Governance",
    "Zielgruppe: große Enterprises mit reifem PMO",
    "Preis-Level: hoch — komplexe Implementierung (12–18 Monate)",
])

# Planisware
add_heading('🎯 Planisware — Leader mit höchster „Completeness of Vision"', level=2)
add_bullets([
    "Referenz-Tool für Engineering-/R&D-Projekte (Stage-Gate, PEP)",
    "Kernstärken: Produktentwicklung, komplexe Ressourcen×Finanz-Verzahnung, Portfolio-Simulationen",
    "Zielgruppe: R&D-getriebene Unternehmen (Automotive, Pharma, Industrie — strukturell nah an STIHL)",
    "Preis-Level: sehr hoch, personell-hungrig — dedizierte Experten nötig (18–24 Monate)",
])

# ServiceNow
add_heading("🔷 ServiceNow SPM Pro — Leader im SPM Magic Quadrant (2024/2025)", level=2)
add_bullets([
    "Teil der Now Platform, native Integration mit ITSM, ITOM, HR, CSM, GRC",
    "Kernstärken: Plattform-Integration, Demand-to-Delivery, App Engine, Now Assist (KI)",
    "Zielgruppe: Enterprises mit vorhandener SNOW-Basis, Fokus auf End-to-End-Digitalisierung",
    "Preis-Level: mittel bis hoch, aber Skaleneffekt bei bestehendem SNOW-Vertrag",
])

doc.add_page_break()

# ============================================================
# VERGLEICHSMATRIX
# ============================================================
add_heading("Head-to-Head: Projekt-Management-Fähigkeiten", level=1)
add_para(
    "Die folgende Matrix vergleicht die drei Player anhand von 16 Kriterien im Kontext "
    "Projekt-Management (5 Sterne = Best-in-Class).",
    italic=True, color=GREY_TXT, space_after=10
)

compare_rows = [
    ("Klassische Projektplanung (Gantt, WBS, Meilensteine)", "★★★★★", "★★★★★", "★★★★"),
    ("Multi-Projekt-Abhängigkeiten",                        "★★★★★", "★★★★★", "★★★★"),
    ("Ressourcen-Management (Skills, Capacity, Optimierung)","★★★★★", "★★★★★", "★★★★"),
    ("Was-wäre-wenn / Szenario-Analysen",                    "★★★★★", "★★★★★", "★★★★ (Pro)"),
    ("Finanz-Management (Kosten, Forecasts, Investment)",    "★★★★★", "★★★★★", "★★★★"),
    ("Programm-Management",                                  "★★★★★", "★★★★★", "★★★★"),
    ("Stage-Gate / PEP-Prozesse",                            "★★★★",  "★★★★★", "★★★ / ★★★★ mit App Engine"),
    ("Agile / Hybrid-PM (SAFe, Scrum)",                     "★★★★★", "★★★★",  "★★★★ (EAP Pro)"),
    ("Reporting & Live-Dashboards",                          "★★★★★", "★★★★",  "★★★★★"),
    ("KI / GenAI (Vorhersagen, Auto-Reports)",              "★★★",   "★★★",   "★★★★★ (Now Assist)"),
    ("Zeiterfassung + Approval",                             "★★★★",  "★★★★★", "★★★★"),
    ("Integration mit ITSM / bestehende SNOW-Landschaft",   "★★",    "★★",    "★★★★★ (nativ)"),
    ("Extensibility / Low-Code Custom",                      "★★★",   "★★★",   "★★★★★ (App Engine)"),
    ("User Experience / Familiarität bei SNOW-Kunden",       "★★",    "★★",    "★★★★"),
    ("Time-to-Value bei bestehender SNOW-Basis",             "★",     "★",     "★★★★★"),
    ("TCO (5J) bei bestehender SNOW-Basis",                  "★★",    "★",     "★★★★"),
]
build_table(
    headers=["Fähigkeit", "Planview", "Planisware", "ServiceNow SPM Pro"],
    rows=compare_rows,
    col_widths_cm=[7.5, 3.0, 3.0, 3.5],
    header_fill=NAVY_HEX,
    center_cols=[1, 2, 3],
)
doc.add_paragraph()

# ============================================================
# FACH-BEWERTUNG
# ============================================================
add_heading("Fach-Bewertung nach Projekt-Management-Aspekten", level=1)

add_heading("Wo Planisware technisch überlegen ist", level=3, color=RED_HL)
add_bullets([
    "Tiefe Stage-Gate-Prozesse für Produktentwicklung (STIHL EWW/PEP wäre hier grundsätzlich sehr gut aufgehoben)",
    "Sehr komplexe Ressourcen×Finanz×Zeit-Verschränkungen über hunderte parallele NPD-Projekte",
    "Referenz-Domäne: Automotive, Aerospace, Pharma R&D",
    "ABER: Insel-Betrieb, kein Anschluss an ITSM/HR/CSM",
])

add_heading("Wo Planview überlegen ist", level=3, color=RED_HL)
add_bullets([
    "Portfolio-Optimierung + Kapazitätsplanung über sehr große PMOs",
    "Skills-basierte Ressourcen-Matching (Best-in-Class)",
    "Adaptive PM (klassisch/agil/hybrid) auf einheitlicher Datenbasis",
    "ABER: komplexe Implementierung, hohe TCO, kein SNOW-Anschluss",
])

add_heading("Wo ServiceNow SPM Pro überlegen ist", level=3, color=GREEN)
add_bullets([
    "Plattform-Integration — Demand aus ITSM fließt direkt in Projekte",
    "Time-to-Value — Team + Betrieb + UI bereits bekannt",
    "KI durchgängig — Now Assist für Reporting, Risikofrüherkennung, Story Generation",
    "Erweiterbarkeit über App Engine — spezifische PEP-Steps ohne 3rd-Party-Tool nachbaubar",
    "Ganzheitliche Sicht über IT-Projekte, Business-Projekte, HR-Projekte, Compliance",
    "ABER: in tiefstem Engineering-PM etwas weniger reif als spezialisierte Tools",
])

doc.add_page_break()

# ============================================================
# EMPFEHLUNG STIHL-KONTEXT
# ============================================================
add_heading("Empfehlung für den STIHL-Kontext", level=1)
add_para(
    "Setup: ServiceNow seit 2014 im Einsatz  ·  SPM Demand bereits aktiv  ·  Migration von altem PM-Tool (PIT/MSPO)",
    italic=True, color=GREY_TXT, space_after=10
)

add_callout(
    "Klare Empfehlung: ServiceNow SPM Pro als strategische PPM-Zielplattform. "
    "Add-on zur bestehenden SNOW-Basis statt Neu-Vendor mit Integrationsaufwand.",
    color_hex="FDE6CC",
    title="👉 Empfehlung"
)

add_heading("Argumente in Vorstands-Reihenfolge", level=2)
argu_rows = [
    ("1", "Datenkontinuität — Demand fließt bereits in SNOW → Projekt-Umsetzung ohne Bruch/Bridge",
         "Sofortiger Prozess-Fluss Demand → Projekt"),
    ("2", "Betriebskontinuität — ITSM-Team seit 12 Jahren im Haus, keine neue Vendor-Beziehung nötig",
         "~ 2 FTE eingespart vs. neuer Vendor"),
    ("3", "UI-Konsistenz — Anwender kennen SNOW-Look, deutlich geringerer Change-Aufwand",
         "~ 200 k€ eingesparte Trainings"),
    ("4", "TCO — Add-on zur bestehenden Plattform statt Zweitvendor + Integration",
         "~ 1,5–2,5 M€/5J. günstiger als Planview/Planisware"),
    ("5", "KI-Roadmap — Now Assist skaliert über alle Bereiche, nicht nur PM",
         "Zukunftsinvestition schützt"),
    ("6", "App Engine für PEP — spezifische EWW-Anforderungen ohne 3rd-Party abbildbar",
         "Kein Doppel-System"),
]
build_table(
    headers=["#", "Argument", "Impact"],
    rows=argu_rows,
    col_widths_cm=[1.0, 10.5, 5.5],
    header_fill=NAVY_HEX,
)
doc.add_paragraph()

add_heading("Zwei ehrliche Vorbehalte (Vorstands-tauglich)", level=2)

add_heading("1. Engineering-PM-Tiefe", level=3, color=NAVY)
add_para(
    "Für die absolute Spitzentiefe im Stage-Gate-PEP (EWW) ist Planisware technisch weiter. "
    "Wenn STIHL in EWW industrialisierte Produktentwicklung mit > 100 parallelen R&D-Projekten "
    "hat und dort maximale Reife will → hybrider Ansatz denkbar:"
)
add_bullets([
    "ServiceNow SPM Pro: Konzern-Portfolio, Demand, Strategy, IT/Operations-PM, Reporting",
    "Planisware (nur EWW): Deep-R&D-PEP-Steuerung mit Übergabe an SNOW für Portfolio-Sicht",
])
add_para(
    "Aber: dieser Split verdoppelt Betriebskosten. Nur sinnvoll, wenn App-Engine-Erweiterung "
    "nachweislich nicht ausreicht.",
    color=RED_HL, italic=True
)

add_heading("2. Sehr komplexe Portfolio-Simulation", level=3, color=NAVY)
add_para(
    'Für „Was-wäre-wenn-Analysen" über 500+ Projekte × Ressourcen × Finanzen simultan hat '
    'Planview noch einen kleinen Vorsprung. Für die typische STIHL-Portfolio-Größe irrelevant.'
)

doc.add_page_break()

# ============================================================
# MIGRATIONS-EMPFEHLUNG
# ============================================================
add_heading("Migrations-Empfehlung von altem PM-Tool", level=1)
add_para(
    'Statt „Big Bang" auf Planview/Planisware — schrittweise Erweiterung ServiceNow SPM Pro:',
    space_after=8
)

phase_rows = [
    ("Phase 0 (heute)", "läuft",       "Demand-Management bereits aktiv"),
    ("Phase 1",         "12 Monate",   "Portfolio + Strategy + Standard-PM (Gantt, Ressourcen, Finanzen) → Ablöse PIT/MSPO für Standard-Projekte"),
    ("Phase 2",         "12 Monate",   "EAP/Agile + tiefes Ressourcen-Management + PEP über App Engine → Ablöse EWW-Spezifika (evtl. mit Planisware-Hedge-Bewertung)"),
    ("Phase 3+",        "rollierend",  "HR-SD, CSM, GRC — SNOW-Plattform breiter nutzen"),
]
build_table(
    headers=["Phase", "Dauer", "Scope"],
    rows=phase_rows,
    col_widths_cm=[3.5, 2.5, 11.0],
    header_fill=ORANGE_HEX,
    center_cols=[1],
)
doc.add_paragraph()

add_callout(
    "Nach Phase 1, wenn sich zeigt, dass App Engine PEP nicht in vernünftiger Tiefe abbildbar "
    "ist — dann Planisware als spezialisiertes Engineering-PM in Betracht ziehen. "
    "Bis dahin: Standard-First mit klarer Migration-Klausel.",
    color_hex="FFF6D5",
    title='⚠ Break-Punkt für „Doch Planisware"-Diskussion'
)

# ============================================================
# FAZIT
# ============================================================
add_heading("Fazit", level=1)
add_para(
    "Für Unternehmen mit bestehender ServiceNow-Basis wie STIHL überwiegen die Vorteile "
    "einer SNOW SPM Pro-Migration klar die Funktionsdefizite gegenüber spezialisierten "
    "PM-Tools. Die entscheidenden Argumente sind nicht in der PM-Funktionstiefe zu finden, "
    "sondern in Plattform-Integration, Betriebskosten, Zeit-bis-Nutzen und der Konsistenz "
    "der Anwenderführung.",
    space_after=8
)
add_para(
    "Planisware bleibt technisch überlegen in tiefsten R&D-Stage-Gate-Prozessen und "
    "sollte als Hedge-Option in der Detail-Auftragsklärung (Vorgang 3 im Vorstandsprotokoll) "
    "bewertet werden. Planview ist technisch stärker in reiner Portfolio-Optimierung, "
    "aber ohne SNOW-Vorteil zu teuer für den STIHL-Kontext.",
    space_after=12
)

# Footer
p = doc.add_paragraph()
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
top = OxmlElement("w:top")
top.set(qn("w:val"), "single")
top.set(qn("w:sz"), "4")
top.set(qn("w:space"), "1")
top.set(qn("w:color"), "D1D5DB")
pBdr.append(top)
pPr.append(pBdr)

p = doc.add_paragraph()
r = p.add_run(
    "Quellen: Gartner Magic Quadrant für Adaptive Project Management and Reporting (Juni 2026); "
    "Gartner Magic Quadrant Strategic Portfolio Management (2024/2025); "
    "Analyse basiert auf öffentlich verfügbaren Vendor-Informationen und Marktbeobachtung. "
    "Bindende Bewertung nur nach Detail-Prüfung im Rahmen des RFP-Prozesses (Vorgang 6, Q4/2026)."
)
r.font.name = "Calibri"
r.font.size = Pt(8.5)
r.font.color.rgb = GREY_TXT
r.italic = True

# Autoren
p = doc.add_paragraph()
r = p.add_run("Autoren: Alex Passaro · Torsten Zahn  ·  Stand: 15.10.2026")
r.font.name = "Calibri"
r.font.size = Pt(8.5)
r.font.color.rgb = GREY_TXT
r.italic = True

doc.save(DST)
print(f"Saved: {DST}")
