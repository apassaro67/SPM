"""Roadmap v11: fix milestone overlap, remove quarter labels & grid lines
inside lanes, improve bar text readability.
- Milestone labels alternate above/below the timeline line.
- Quarter band removed; year bar alone with subtle quarter ticks under year.
- Inside lanes: no vertical grid, larger/clearer text on bars.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v9.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v11.pptx"

ORANGE   = RGBColor(0xF0, 0x7F, 0x12)
ORANGE_LT= RGBColor(0xFB, 0xD3, 0xA8)
ORANGE2  = RGBColor(0xE9, 0x6C, 0x0C)
ORANGE3  = RGBColor(0xC9, 0x53, 0x00)
DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
LIGHT_BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
GREEN    = RGBColor(0x2E, 0x7D, 0x32)
BLUE_HD  = RGBColor(0x3B, 0x6B, 0xA5)
RED_HL   = RGBColor(0xB9, 0x1C, 0x1C)
GREY_LN  = RGBColor(0xCB, 0xD2, 0xD9)
LIGHT_GREY = RGBColor(0xD1, 0xD5, 0xDB)

U = 76200
def ex(u): return int(u * U)

prs = Presentation(SRC)
slide = prs.slides[13]

for shp in list(slide.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

def add_rect(left, top, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
    s = slide.shapes.add_shape(shape, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def add_diamond(left, top, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, ex(left), ex(top), ex(w), ex(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.0)
    s.shadow.inherit = False
    return s

def set_paragraphs(shape, paragraphs, *, default_size=10, default_color=DARK, default_align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(40000); tf.margin_right = Emu(40000)
    tf.margin_top = Emu(15000); tf.margin_bottom = Emu(15000)
    tf.clear()
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, st = item, {}
        else:
            text, st = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = st.get("align", default_align)
        run = p.add_run(); run.text = text
        run.font.name = "Calibri"
        run.font.size = Pt(st.get("size", default_size))
        run.font.bold = st.get("bold", False)
        run.font.italic = st.get("italic", False)
        run.font.color.rgb = st.get("color", default_color)
        if "space_after" in st:
            p.space_after = Pt(st["space_after"])

def add_text(left, top, w, h, paragraphs, *, anchor=MSO_ANCHOR.MIDDLE, **kw):
    tb = slide.shapes.add_textbox(ex(left), ex(top), ex(w), ex(h))
    tb.fill.background(); tb.line.fill.background()
    tb.text_frame.vertical_anchor = anchor
    set_paragraphs(tb, paragraphs, **kw)
    return tb

# ========== HEADER ==========
add_rect(0, 0, 160, 6, DARK)
add_text(3, 0.4, 140, 2.6, [
    ("Roadmap 2026 – 2028:  Konzeption · Umsetzung · Go-Live",
     dict(size=20, bold=True, color=WHITE)),
], anchor=MSO_ANCHOR.TOP)
add_text(3, 3.2, 140, 2.4, [
    ("Drei-Jahres-Plan SNOW SPM mit Phasen, Meilensteinen und vertraglichen Eckpunkten",
     dict(size=12, italic=True, color=LIGHT_GREY)),
], anchor=MSO_ANCHOR.TOP)
add_rect(150, 0, 10, 6, ORANGE)
add_text(150, 0, 10, 6, [
    ("Roadmap", dict(size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
])

# ========== TIMELINE GEOMETRY ==========
LANE_LBL_X = 4
LANE_LBL_W = 26
TL_X = LANE_LBL_X + LANE_LBL_W + 1.2
TL_W = 156 - TL_X
N_QUARTERS = 12
QW = TL_W / N_QUARTERS
def qx(qi): return TL_X + qi * QW

# ========== YEAR AXIS ONLY (Q-Beschriftung entfernt) ==========
axis_top = 7.5
year_h = 3.5
years = [
    (2026, RGBColor(0xF7, 0xC8, 0x90)),
    (2027, RGBColor(0xF4, 0xA4, 0x4C)),
    (2028, RGBColor(0xE0, 0x82, 0x1B)),
]
for i, (y, col) in enumerate(years):
    x = qx(i * 4)
    add_rect(x, axis_top, 4 * QW - 0.2, year_h, col,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(x, axis_top, 4 * QW - 0.2, year_h, [
        (str(y), dict(size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)),
    ], anchor=MSO_ANCHOR.MIDDLE)

# Subtle quarter ticks under the year bar (no labels)
tick_y = axis_top + year_h
for i in range(1, N_QUARTERS):
    if i % 4 == 0:
        continue  # year boundaries already separated by gap
    add_rect(qx(i) - 0.05, tick_y, 0.08, 0.6, GREY_LN)

add_text(LANE_LBL_X, axis_top, LANE_LBL_W, year_h, [
    ("Zeit ▶", dict(size=11, bold=True, color=GREY_TXT, align=PP_ALIGN.RIGHT)),
], anchor=MSO_ANCHOR.MIDDLE)

# ========== MILESTONES ROW (alternating up/down) ==========
ms_top = axis_top + year_h + 1.8
ms_h = 11.0   # taller to accommodate above/below labels
line_y = ms_top + ms_h / 2

def qpos_for(year, month):
    base = (year - 2026) * 4
    qi = (month - 1) // 3
    frac = ((month - 1) % 3 + 1) / 3.0
    return base + qi + frac - 0.5

milestones = [
    (qpos_for(2026, 4),  "PPB 08.04",         "Projekt-Board",            BLUE_HD, False, "below"),
    (qpos_for(2026, 5),  "VS 18.05",          "heutige Vorstandssitzung", RED_HL,  True,  "above"),
    (qpos_for(2026, 9),  "Jahresplanung",     "01.09.2026",               BLUE_HD, False, "below"),
    (qpos_for(2026, 11), "Kick-off Phase 1",  "Q3/Q4 2026",               ORANGE,  False, "above"),
    (qpos_for(2027, 12), "Go-Live Phase 1",   "Org + IT Projekte",        GREEN,   True,  "below"),
    (qpos_for(2028, 12), "Go-Live Phase 2",   "VEW / VPM",                GREEN,   True,  "above"),
]

add_text(LANE_LBL_X, ms_top, LANE_LBL_W, ms_h, [
    ("MEILENSTEINE", dict(size=11, bold=True, color=DARK, align=PP_ALIGN.RIGHT)),
], anchor=MSO_ANCHOR.MIDDLE)

# Horizontal line
add_rect(TL_X, line_y - 0.05, TL_W, 0.15, GREY_LN)

LBL_W = 18  # label box width
for q_pos, head, sub, col, big, side in milestones:
    cx = qx(q_pos)
    diam = 2.2 if big else 1.6
    add_diamond(cx - diam / 2, line_y - diam / 2, diam, diam, col)
    # Leader line
    if side == "above":
        # leader from diamond up
        leader_y0 = line_y - diam / 2 - 0.1
        leader_y1 = ms_top + 0.5
        add_rect(cx - 0.04, leader_y1, 0.08, leader_y0 - leader_y1, GREY_LN)
        # label box above
        add_text(cx - LBL_W / 2, ms_top - 0.6, LBL_W, 2.0, [
            (head, dict(size=9.5, bold=True, color=col, align=PP_ALIGN.CENTER)),
        ], anchor=MSO_ANCHOR.BOTTOM)
        add_text(cx - LBL_W / 2, ms_top + 1.2, LBL_W, 1.6, [
            (sub, dict(size=8, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER)),
        ], anchor=MSO_ANCHOR.MIDDLE)
    else:
        leader_y0 = line_y + diam / 2 + 0.1
        leader_y1 = ms_top + ms_h - 0.5
        add_rect(cx - 0.04, leader_y0, 0.08, leader_y1 - leader_y0, GREY_LN)
        add_text(cx - LBL_W / 2, ms_top + ms_h - 2.6, LBL_W, 1.8, [
            (head, dict(size=9.5, bold=True, color=col, align=PP_ALIGN.CENTER)),
        ], anchor=MSO_ANCHOR.BOTTOM)
        add_text(cx - LBL_W / 2, ms_top + ms_h - 0.7, LBL_W, 1.6, [
            (sub, dict(size=8, italic=True, color=GREY_TXT, align=PP_ALIGN.CENTER)),
        ], anchor=MSO_ANCHOR.MIDDLE)

# ========== SWIM LANES ==========
lane_top = ms_top + ms_h + 1.0
lane_h = 9.0
lane_gap = 1.0

def draw_lane(idx, label, sublabel, lbl_color, bars):
    y = lane_top + idx * (lane_h + lane_gap)
    add_rect(LANE_LBL_X, y, LANE_LBL_W, lane_h, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_rect(LANE_LBL_X, y, 0.8, lane_h, lbl_color)
    add_text(LANE_LBL_X + 1.5, y + 0.5, LANE_LBL_W - 2, 3.0, [
        (label, dict(size=11.5, bold=True, color=DARK)),
    ], anchor=MSO_ANCHOR.TOP)
    add_text(LANE_LBL_X + 1.5, y + 3.4, LANE_LBL_W - 2, lane_h - 4, [
        (sublabel, dict(size=9, italic=True, color=GREY_TXT)),
    ], anchor=MSO_ANCHOR.TOP)
    # Lane background (no internal grid)
    add_rect(TL_X, y, TL_W, lane_h, LIGHT_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    # Subtle year boundaries inside lane
    for yi in (1, 2):
        add_rect(qx(yi * 4) - 0.06, y + 0.4, 0.12, lane_h - 0.8, GREY_LN)
    # Bars
    for q_start, q_end, color, head, sub in bars:
        x0 = qx(q_start)
        w  = qx(q_end) - x0 - 0.2
        bar_y = y + 1.1
        bar_h = lane_h - 2.2
        add_rect(x0 + 0.1, bar_y, w, bar_h, color,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        # Pick font sizes and split based on bar width
        title_size = 11 if w >= 18 else (10 if w >= 12 else 9)
        sub_size   = 9.5 if w >= 18 else (8.5 if w >= 12 else 7.5)
        if sub:
            add_text(x0 + 0.6, bar_y + 0.2, w - 1.2, bar_h * 0.55, [
                (head, dict(size=title_size, bold=True, color=WHITE)),
            ], anchor=MSO_ANCHOR.BOTTOM)
            add_text(x0 + 0.6, bar_y + bar_h * 0.50, w - 1.2, bar_h * 0.5, [
                (sub, dict(size=sub_size, italic=True, color=WHITE)),
            ], anchor=MSO_ANCHOR.TOP)
        else:
            add_text(x0 + 0.6, bar_y, w - 1.2, bar_h, [
                (head, dict(size=title_size, bold=True, color=WHITE)),
            ], anchor=MSO_ANCHOR.MIDDLE)

# Lane content (text shortened for narrow bars)
lane_bars_1 = [
    (0.2, 2.8, ORANGE_LT, "Auftragsklärung · Stakeholder · Phasen-/Kostenplanung", "Q1 – Q3 2026"),
    (2.4, 4.2, ORANGE,    "Partnerauswahl SNOW SPM",                                "Q3/Q4 2026"),
    (3.4, 5.0, ORANGE2,   "SNOW Budgetierung & Contract Renewal",                   "ab Q4 2026"),
]
draw_lane(0, "Vorbereitung & Vertrag", "Auftragsklärung · Budget · Partnerauswahl",
          BLUE_HD, lane_bars_1)

lane_bars_2 = [
    (3.0, 4.0, ORANGE_LT, "Konzeption", "Demand · Strategy · MDM · Portfolio"),
    (4.0, 8.0, ORANGE,    "Umsetzung Phase 1", "+ PIT-Schnittstelle"),
    (7.6, 8.0, GREEN,     "Go-Live", "Org + IT"),
]
draw_lane(1, "Phase 1  SNOW SPM", "Demand · Strategy · MDM · Portfolio · Workflows",
          ORANGE, lane_bars_2)

lane_bars_3 = [
    (6.0, 8.0, ORANGE_LT, "Konzeption Phase 2", "Project Mgmt · PEP · VPM"),
    (8.0, 11.7, ORANGE2,  "Umsetzung Phase 2",  "+ Maßnahmen-Mgmt"),
    (11.6, 12.0, GREEN,   "Go-Live", "VEW / VPM"),
]
draw_lane(2, "Phase 2  SNOW SPM", "Project Management · PEP · VPM · Portfolio",
          ORANGE3, lane_bars_3)

lane_bars_4 = [
    (3.0, 6.0, NAVY,
     "Alternative PM-Plattform (VEW / VPM)",
     "Bewertung als Risiko-Hedge"),
]
draw_lane(3, "Alternative-Spur", "PM-Plattform für VEW / VPM (Hedge)",
          NAVY, lane_bars_4)

# ========== LEGEND ==========
leg_top = lane_top + 4 * (lane_h + lane_gap) + 0.3
leg_items = [
    ("Konzeption",   ORANGE_LT),
    ("Umsetzung",    ORANGE),
    ("Umsetzung +",  ORANGE2),
    ("Go-Live",      GREEN),
    ("Vorbereitung / Vertrag", BLUE_HD),
    ("Alternative",  NAVY),
]
lx = LANE_LBL_X
add_text(lx, leg_top, 22, 2.0, [
    ("Legende:", dict(size=9.5, bold=True, color=DARK)),
], anchor=MSO_ANCHOR.MIDDLE)
cursor = lx + 11
for label, col in leg_items:
    add_rect(cursor, leg_top + 0.5, 1.6, 1.6, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(cursor + 1.9, leg_top, 18, 2.2, [
        (label, dict(size=8.8, color=DARK)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    cursor += 21

# Vorstands-Meilenstein-Symbol
mleg_x = 130
add_diamond(mleg_x, leg_top + 0.5, 1.6, 1.6, RED_HL)
add_text(mleg_x + 2.0, leg_top, 26, 2.2, [
    ("◆ Vorstands-Meilenstein", dict(size=8.8, color=DARK)),
], anchor=MSO_ANCHOR.MIDDLE)

# ========== FOOTER ==========
add_text(4, 86.3, 152, 2.2, [
    ("Konsolidiert aus den Original-Folien 2026 / 2027 / 2028 (Rough Schedule).  Stand: 18.05.2026.",
     dict(size=8.5, italic=True, color=GREY_TXT)),
])
add_rect(0, 89.3, 160, 0.7, ORANGE)

# ========== Delete the year slides 15 & 16 ==========
sldIdLst = prs.slides._sldIdLst
slides = list(sldIdLst)
for sld in [slides[14], slides[15]]:
    rId = sld.get(qn("r:id"))
    sldIdLst.remove(sld)
    prs.part.drop_rel(rId)

prs.save(DST)
print("Saved:", DST)
