"""Slide 10 v18: replace donut/percent charts with native traffic-light
(Ampel) indicators. Five cells stay green (≥ 80 %) and Finanzen goes
yellow (60 %, 'bedingt machbar')."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC = "/root/.claude/uploads/169d168f-c9e7-43eb-9b92-6f80b02d3b1f/d655d9ff-2026_05_18_VS_Zielbild_SNOW_SPM_v16.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v18.pptx"

DARK     = RGBColor(0x1F, 0x29, 0x37)
GREY_TXT = RGBColor(0x4B, 0x55, 0x63)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
NAVY     = RGBColor(0x2C, 0x3E, 0x50)
NAVY_DK  = RGBColor(0x18, 0x25, 0x34)
RED      = RGBColor(0xC8, 0x2E, 0x2E)
RED_DIM  = RGBColor(0x4A, 0x2B, 0x2B)
YEL      = RGBColor(0xF5, 0xC2, 0x14)
YEL_DIM  = RGBColor(0x55, 0x49, 0x18)
GRN      = RGBColor(0x2E, 0x9B, 0x32)
GRN_DIM  = RGBColor(0x24, 0x42, 0x28)

prs = Presentation(SRC)
slide = prs.slides[9]
spTree = slide.shapes._spTree

# Find the chart shapes and percent overlays.
# Each cell has Chart 3 (W=1.30 x 1.30) followed by Textplatzhalter 5 (W=0.60 x 0.33 with %).
to_remove = []
percent_cells = []  # list of (left, top, percent_value)
shapes_list = list(slide.shapes)
i = 0
while i < len(shapes_list):
    sh = shapes_list[i]
    if sh.name.startswith("Chart") and abs(sh.width/914400 - 1.30) < 0.05:
        # Get accompanying % textbox
        pct_value = None
        pct_shape = None
        if i + 1 < len(shapes_list):
            nxt = shapes_list[i + 1]
            if nxt.has_text_frame and "%" in nxt.text_frame.text:
                pct_value = int(nxt.text_frame.text.replace("%", "").strip())
                pct_shape = nxt
        percent_cells.append({
            "chart_left": sh.left / 914400,
            "chart_top":  sh.top / 914400,
            "chart_w":    sh.width / 914400,
            "chart_h":    sh.height / 914400,
            "pct": pct_value if pct_value is not None else 80,
        })
        to_remove.append(sh)
        if pct_shape is not None:
            to_remove.append(pct_shape)
            i += 2
            continue
    i += 1

print(f"Cells found: {len(percent_cells)}")
for c in percent_cells:
    print(f"  {c}")

# Remove old chart/percent shapes
for sh in to_remove:
    sp = sh._element
    sp.getparent().remove(sp)

# Helper to add native shapes using inch coordinates (since the original slide
# uses inch-based EMU positions for these elements)
INCH = 914400
def add_rect_in(left_in, top_in, w_in, h_in, fill, line=None,
                shape=MSO_SHAPE.RECTANGLE, line_w=1.0):
    s = slide.shapes.add_shape(shape,
                               int(left_in * INCH), int(top_in * INCH),
                               int(w_in * INCH), int(h_in * INCH))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s

def add_oval_in(left_in, top_in, w_in, h_in, fill, line=None, line_w=0.0):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                               int(left_in * INCH), int(top_in * INCH),
                               int(w_in * INCH), int(h_in * INCH))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    return s

def add_text_in(left_in, top_in, w_in, h_in, text, *, size=10, bold=False,
                italic=False, color=DARK, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(int(left_in * INCH), int(top_in * INCH),
                                  int(w_in * INCH), int(h_in * INCH))
    tb.fill.background(); tb.line.fill.background()
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(20000); tf.margin_right = Emu(20000)
    tf.margin_top = Emu(10000); tf.margin_bottom = Emu(10000)
    tf.vertical_anchor = anchor
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = "Calibri"; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return tb

def ampel(left, top, pct):
    """Draw a vertical Ampel + % value + status label,
    centered horizontally in a 1.3" wide area starting at (left, top)."""
    # Status mapping
    if pct >= 80:
        status, status_color = "Machbar", GRN
        on = "green"
    elif pct >= 50:
        status, status_color = "Bedingt machbar", YEL
        on = "yellow"
    else:
        status, status_color = "Nicht machbar", RED
        on = "red"

    # Casing dimensions — vertical traffic light
    case_w = 0.55
    case_h = 1.20
    case_x = left + (1.30 - case_w) / 2  # center within 1.3" chart slot
    case_y = top - 0.05  # slightly above original chart top to leave room for label
    add_rect_in(case_x, case_y, case_w, case_h, NAVY_DK,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)

    # 3 lights
    light_d = 0.30
    pad = (case_h - 3 * light_d) / 4  # vertical padding between lights
    light_x = case_x + (case_w - light_d) / 2
    lights = [
        ("red",    RED if on == "red" else RED_DIM,        on == "red"),
        ("yellow", YEL if on == "yellow" else YEL_DIM,     on == "yellow"),
        ("green",  GRN if on == "green" else GRN_DIM,      on == "green"),
    ]
    for i, (_name, color, is_on) in enumerate(lights):
        ly = case_y + pad + i * (light_d + pad)
        # Highlight ring on the active one
        if is_on:
            add_oval_in(light_x - 0.02, ly - 0.02,
                        light_d + 0.04, light_d + 0.04, color)
        add_oval_in(light_x, ly, light_d, light_d, color)

    # Status label + percent under the casing
    label_y = case_y + case_h + 0.04
    label_w = 1.30
    label_x = left
    add_text_in(label_x, label_y, label_w, 0.22,
                f"{pct}%", size=11, bold=True, color=status_color,
                align=PP_ALIGN.CENTER)
    add_text_in(label_x, label_y + 0.22, label_w, 0.22,
                status, size=9, italic=True, color=GREY_TXT,
                align=PP_ALIGN.CENTER)

for cell in percent_cells:
    ampel(cell["chart_left"], cell["chart_top"], cell["pct"])

prs.save(DST)
print("Saved:", DST)
