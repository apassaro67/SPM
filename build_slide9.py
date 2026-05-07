"""Fill slide 9 (Vorprojekt SNOW – Ergebnisse 1 von 2) with the
team-statement content from the workshop image.

Mapping per cell (matching the existing slide layout):
  Top row     | Entwicklung (EWW)        | Produktion (TGG)        | IT (VFI)
  Bottom row  | Strategie (SCD)          | Produktmanagement (MPM) | Finanzen (VFI)
"""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v2.pptx"
DST = "/tmp/2026_05_18_VS_Zielbild_SNOW_SPM_v3.pptx"

DARK = RGBColor(0x1F, 0x29, 0x37)

# 3 bullets per cell (kept short for board readability)
content = {
    # idx in shape list -> list of bullet strings
    21: [  # Entwicklung (EWW) – Produktprojekte – 80% (unter Umständen)
        "Szenario grundsätzlich abbildbar – Thumbs Up für Auftragsklärung",
        "PM-Anforderungen tiefer legen, Schwerpunkt Ressourcen- und Finanzplanung",
        "Customizing über Low-Code hinaus erwartet; Anbietervergleich/Benchmarking ergänzen",
    ],
    22: [  # Produktion (TGG) – Infrastrukturprojekte – 80%
        "Szenario grundsätzlich abbildbar – Thumbs Up für Auftragsklärung",
        "Anforderungen müssen inhaltlich noch tiefergelegt werden",
        "Sehr guter Implementierungspartner nötig; Prozesse vorab klar definieren",
    ],
    23: [  # IT (VFI) – IT-Projekte – 80%
        "IT-Szenario und FB-Anforderungen grundsätzlich abbildbar – Thumbs Up",
        "Deutliches Potenzial der strategischen Plattform ServiceNow für die Use Cases",
        "Detaillierte Auftragsklärung rechtzeitig zur Jahresplanung weiterverfolgen",
    ],
    24: [  # Strategie (SCD) – Sonderprojekte – 80%
        "SPM-Szenarien (Portfolio Management & OKR) grundsätzlich abbildbar – Thumbs Up",
        "Mächtigkeit der Plattform verlangt sauberes Scope- und Rollout-Management",
        "Potenzial: zentrale Datenverfügbarkeit; Zusatzaufwand in Prozess-/Verantwortungsklärung",
    ],
    25: [  # Produktmanagement (MPM) – Strategische Produktplanung – 80%
        "SPP Gantt und Bottom-up Planung auf Modell-Ebene grundsätzlich abbildbar – Thumbs Up",
        "Anforderungsmanagement integrierbar (LH/CRD → Steckbrief): zusätzliches Potenzial",
        "Chance: durchgängiger Systemansatz Strategie → Idee → Vorhaben → Projekt",
    ],
    26: [  # Finanzen (VFI) – EV-Prozess – 60%
        "Szenario abbildbar – Thumbs Up für Auftragsklärung",
        "Viele Anforderungen müssen inhaltlich noch tiefergelegt werden",
        "Detail-Klärung in Folge-Workshops zur weiteren Schärfung",
    ],
}

# Index of the percent text overlays (Textplatzhalter 5)
percent_updates = {
    28: "80%",  # Entwicklung
    30: "80%",  # Produktion
    32: "80%",  # IT
    34: "80%",  # Strategie
    36: "80%",  # Produktmanagement
    38: "60%",  # Finanzen
}

prs = Presentation(SRC)
slide = prs.slides[8]

shapes = list(slide.shapes)

def set_bullets(shape, bullets, *, size=10):
    tf = shape.text_frame
    tf.word_wrap = True
    # Reset
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        run = p.add_run()
        run.text = b
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.color.rgb = DARK

for idx, bullets in content.items():
    set_bullets(shapes[idx], bullets, size=10)

# Update percent labels
for idx, val in percent_updates.items():
    sh = shapes[idx]
    tf = sh.text_frame
    # Preserve existing formatting; just swap text in first run
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = val
    else:
        tf.text = val

prs.save(DST)
print("Saved:", DST)
